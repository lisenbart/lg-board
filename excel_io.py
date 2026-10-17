"""Read/write SP_MGX_check.xlsx as a calculator: hours stay = shifts × 9."""
from __future__ import annotations

import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
XLSX_PATH = ROOT / "data" / "SP_MGX_check.xlsx"
NAME_ID = re.compile(r"^([A-Z]{2}-[A-Z0-9]+-\d+(?:-\d+)?)", re.I)
LG_ID = re.compile(r"lg-id:([A-Za-z0-9_-]+)")

HEADER_FILL = PatternFill("solid", fgColor="D8D8D8")
DX_FILL = PatternFill("solid", fgColor="9DC3E6")
DD_FILL = PatternFill("solid", fgColor="BFBFBF")
DS_FILL = PatternFill("solid", fgColor="C6B19B")
WIP_FILL = PatternFill("solid", fgColor="59BBFF")
CLOSED_FILL = PatternFill("solid", fgColor="92D050")
THIN = Border(
    left=Side(style="hair", color="B4B4B4"),
    right=Side(style="hair", color="B4B4B4"),
    top=Side(style="hair", color="B4B4B4"),
    bottom=Side(style="hair", color="B4B4B4"),
)
CALIBRI = Font(name="Calibri", size=11)
CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
CALIBRI_TOTAL = Font(name="Calibri", size=12, bold=True, color="083358")
TOTAL_FILL = PatternFill("solid", fgColor="579BFC")
CENTER = Alignment(horizontal="center", vertical="center")

LINE_FILL = {"DX": DX_FILL, "DD": DD_FILL, "DS": DS_FILL}
LINE_MARKS = {"DX", "DD", "DS"}
ROW_LABELS = {"mgx", "shifts", "hours", "deadline", "project name", "d.line", "empl"}
PROJECT_NAME_RE = re.compile(r"^(DD|DX|DS)-.+$", re.I)

# ECMA theme index: 0 lt1, 1 dk1, 2 lt2, 3 dk2, 4–9 accents. Google Sheets
# paints WIP blue as theme 8 (accent5) and Closed green as theme 9 (accent6).
_THEME_FALLBACK = [
    "FFFFFF",
    "000000",
    "E7E6E6",
    "44546A",
    "5B9BD5",
    "ED7D31",
    "A5A5A5",
    "FFC000",
    "5B9BD5",
    "70AD47",
    "0563C1",
    "954F72",
]
_THEME_NAMES = (
    "lt1",
    "dk1",
    "lt2",
    "dk2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
)
_WIP_STATUSES = {
    "wip",
    "revision",
    "done",
    "approve",
    "new",
    "storyboard",
    "clientart",
    "finalart",
    "build",
    "htmlready",
}


def _apply_tint(rgb: str, tint: float) -> str:
    raw = str(rgb or "").replace("#", "").upper()
    if len(raw) != 6:
        return raw
    try:
        t = float(tint or 0)
    except (TypeError, ValueError):
        return raw
    if abs(t) < 1e-9:
        return raw
    out = []
    for i in range(0, 6, 2):
        c = int(raw[i : i + 2], 16)
        if t < 0:
            c = int(round(c * (1 + t)))
        else:
            c = int(round(c + (255 - c) * t))
        out.append(f"{max(0, min(255, c)):02X}")
    return "".join(out)


def theme_rgbs(wb) -> list[str]:
    """Resolve workbook theme scheme colors to RRGGBB, else Office-ish fallback."""
    colors = list(_THEME_FALLBACK)
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return colors
    try:
        from xml.etree import ElementTree as ET

        root = ET.fromstring(raw)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = root.find(".//a:clrScheme", ns)
        if scheme is None:
            return colors
        by_name = {}
        for child in list(scheme):
            tag = child.tag.rsplit("}", 1)[-1]
            srgb = child.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
            sys_el = child.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}sysClr")
            val = ""
            if srgb is not None:
                val = str(srgb.get("val") or "")
            elif sys_el is not None:
                val = str(sys_el.get("lastClr") or "")
            val = val.replace("#", "").upper()
            if len(val) == 8:
                val = val[2:]
            if len(val) == 6:
                by_name[tag] = val
        for i, name in enumerate(_THEME_NAMES):
            if by_name.get(name):
                colors[i] = by_name[name]
    except Exception:
        return list(_THEME_FALLBACK)
    return colors


def _cell_fill_hex(cell, theme: list[str] | None = None) -> str:
    fill = getattr(cell, "fill", None)
    if not fill or fill.patternType in (None, "none"):
        return ""
    color = fill.fgColor
    if not color:
        return ""
    ctype = getattr(color, "type", None)
    tint = getattr(color, "tint", 0) or 0
    palette = theme or _THEME_FALLBACK
    if ctype == "theme":
        try:
            idx = int(color.theme)
        except (TypeError, ValueError):
            return ""
        if 0 <= idx < len(palette):
            return _apply_tint(palette[idx], tint)
        return ""
    if ctype == "rgb" and color.rgb:
        raw = str(color.rgb).replace("#", "").upper()
        if raw in {"00000000", "FFFFFFFF"}:
            return ""
        if len(raw) == 8:
            raw = raw[2:]
        return _apply_tint(raw, tint) if len(raw) == 6 else ""
    return ""


def status_from_line_fill(hex_color: str) -> str:
    """Left DX/DD cell: green = Closed, blue = WIP. Same marks as SP_MGX_check."""
    raw = str(hex_color or "").replace("#", "").upper()
    if len(raw) == 8:
        raw = raw[2:]
    if len(raw) != 6:
        return ""
    r, g, b = int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)
    if g >= 150 and g > b + 20 and g >= r:
        return "closed"
    if b >= 150 and b > r + 20:
        return "wip"
    return ""


def line_fill_for_status(status: str) -> PatternFill | None:
    if status == "closed":
        return CLOSED_FILL
    if status in _WIP_STATUSES:
        return WIP_FILL
    return None


def _hex_fill(color: str) -> PatternFill:
    raw = (color or "92D050").replace("#", "")
    if len(raw) == 8:
        raw = raw[2:]
    return PatternFill("solid", fgColor=raw[:6] or "92D050")


def sheet_name_for(ym: str) -> str:
    y, m = [int(x) for x in ym.split("-")]
    return f"{datetime(y, m, 1).strftime('%B')} {str(y)[2:]}"


def _sheet_aliases(ym: str) -> set[str]:
    y, m = [int(x) for x in ym.split("-")]
    dt = datetime(y, m, 1)
    yy = str(y)[2:]
    return {
        f"{dt.strftime('%B')} {yy}".lower(),
        f"{dt.strftime('%B')} {y}".lower(),
        f"{dt.strftime('%b')} {yy}".lower(),
        f"{dt.strftime('%b')} {y}".lower(),
    }


def find_sheet(wb, ym: str):
    want = _sheet_aliases(ym)
    for name in wb.sheetnames:
        if name.strip().lower() in want:
            return wb[name]
    return None


def _looks_like_project(name: str) -> bool:
    return bool(PROJECT_NAME_RE.match(str(name or "").strip()))


def _cell_text(value) -> str:
    if value in (None, ""):
        return ""
    return str(value).strip()


def _project_name(col_a: str, col_b: str) -> str:
    """PROJECT NAME is column B. A project typed only in A still counts."""
    b = col_b.strip()
    a = col_a.strip()
    if b and b.lower() not in ROW_LABELS and b.upper() not in LINE_MARKS:
        return b
    if _looks_like_project(a):
        return a
    return ""


def designers(db: dict) -> list[dict]:
    people = [p for p in db.get("people") or [] if p.get("role") in {"teamlead", "designer"}]
    return sorted(people, key=lambda p: (p.get("payOrder") or 99, p.get("name") or ""))


def _as_shift(value) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, str) and value.strip().startswith("="):
        return 0.0
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def _hours_value(value, shifts: float) -> float:
    if isinstance(value, str) and value.strip().startswith("="):
        return round(shifts * 9, 1)
    try:
        if value not in (None, ""):
            return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        pass
    return round(shifts * 9, 1) if shifts else 0.0


def round_hours(value) -> float:
    try:
        n = float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0
    if n <= 0:
        return 0.0
    return round(n, 1)


def help_entries(task: dict) -> list[dict]:
    """Studio hours given to helpers. Client HOURS stay on the task."""
    who = str(task.get("assigneeId") or "").strip()
    out: list[dict] = []
    seen: set[str] = set()
    for row in task.get("help") or []:
        pid = str((row or {}).get("id") or "").strip()
        hours = round_hours((row or {}).get("hours"))
        if not pid or pid == who or hours <= 0 or pid in seen:
            continue
        seen.add(pid)
        out.append({"id": pid, "hours": hours})
    return out


def client_hours(task: dict, shift_hours: float = 9) -> float:
    try:
        h = float(task.get("hours") or 0)
        if h:
            return round(h, 1)
    except (TypeError, ValueError):
        pass
    try:
        return round(float(task.get("shifts") or 0) * float(shift_hours or 9), 1)
    except (TypeError, ValueError):
        return 0.0


def hours_for_person(task: dict, person_id: str, shift_hours: float = 9) -> float:
    pid = str(person_id or "").strip()
    if not pid:
        return 0.0
    dept = str(task.get("department") or task.get("direction") or "").lower()
    if "play" in dept:
        import playables_io

        return playables_io.hours_for_person(task, pid)
    total = client_hours(task, shift_hours)
    help_rows = help_entries(task)
    given = round(sum(row["hours"] for row in help_rows), 1)
    remainder = max(0.0, round(total - given, 1))
    if pid == str(task.get("assigneeId") or "").strip():
        return remainder
    for row in help_rows:
        if row["id"] == pid:
            return row["hours"]
    return 0.0


def normalize_help(task: dict, people_ids: set[str] | None = None) -> list[dict]:
    who = str(task.get("assigneeId") or "").strip()
    total = client_hours(task)
    rows = []
    used = 0.0
    for row in help_entries(task):
        pid = row["id"]
        if people_ids is not None and pid not in people_ids:
            continue
        if pid == who:
            continue
        hours = row["hours"]
        if used + hours > total + 0.001:
            hours = round(max(0.0, total - used), 1)
        if hours <= 0:
            continue
        rows.append({"id": pid, "hours": hours})
        used = round(used + hours, 1)
        if used >= total:
            break
    task["help"] = rows
    return rows


def _norm_name(name: str) -> str:
    return " ".join(str(name or "").strip().split()).lower()


def people_by_name(people: list[dict]) -> dict:
    """Map visible names and aliases (Маша → Марія) to the crew record."""
    out: dict = {}
    for person in people or []:
        keys = [person.get("name"), *(person.get("aliases") or [])]
        for key in keys:
            n = _norm_name(key)
            if n:
                out[n] = person
    return out


def _id_key(name: str) -> str:
    m = NAME_ID.match(str(name or "").strip())
    return m.group(1).upper() if m else ""


def _deadline_iso(value, ym: str) -> str:
    """Studio dates are day-first. Excel often stores 07/10 as 10 July."""
    y = m = 0
    try:
        y, m = [int(x) for x in str(ym or "").split("-")[:2]]
    except (TypeError, ValueError):
        pass
    if isinstance(value, datetime):
        dt = value
        if m and dt.month != m and dt.day == m:
            try:
                dt = dt.replace(year=y or dt.year, month=m, day=dt.month)
            except ValueError:
                pass
        return dt.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d/%m", "%d.%m"):
        try:
            dt = datetime.strptime(text, fmt)
            if fmt in {"%d/%m", "%d.%m"}:
                dt = dt.replace(year=y or datetime.utcnow().year)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def _person_id_from_cell(cell) -> str:
    note = getattr(cell.comment, "text", "") or ""
    m = LG_ID.search(note)
    return m.group(1) if m else ""


def _stamp_person_header(cell, person: dict) -> None:
    cell.value = person.get("name") or cell.value
    cell.font = CALIBRI_BOLD
    cell.alignment = CENTER
    cell.comment = Comment(f"lg-id:{person['id']}", "LG")


def _header_columns(ws) -> list[dict]:
    cols = []
    for col in range(6, 48):
        cell = ws.cell(1, col)
        raw = cell.value
        pid = _person_id_from_cell(cell)
        if (raw is None or str(raw).strip() == "") and not pid:
            if cols:
                break
            continue
        cols.append(
            {
                "col": col,
                "name": str(raw or "").strip(),
                "id": pid,
            }
        )
    return cols


def parse_workbook(path: Path, ym: str) -> dict:
    empty = {"people": [], "rows": [], "sheet": None}
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as e:
        print(f"[excel] parse {path}: {e}", flush=True)
        return empty
    try:
        ws = find_sheet(wb, ym)
        if ws is None:
            return empty
        theme = theme_rgbs(wb)
        headers = _header_columns(ws)
        people = [h["name"] for h in headers if h["name"]]
        rows = []
        line = "DX"
        for r in range(2, (ws.max_row or 1) + 1):
            a = _cell_text(ws.cell(r, 1).value)
            b = _cell_text(ws.cell(r, 2).value)
            if a.lower() == "total":
                continue
            if a.upper() in LINE_MARKS and (not b or b.lower() in ROW_LABELS):
                line = a.upper()
                continue
            name = _project_name(a, b)
            if not name:
                continue
            shifts_n = _as_shift(ws.cell(r, 3).value)
            hours_n = _hours_value(ws.cell(r, 4).value, shifts_n)
            deadline = ws.cell(r, 5).value
            shares = []
            for h in headers:
                val = ws.cell(r, h["col"]).value
                if val in (None, ""):
                    continue
                hh = _hours_value(val, shifts_n)
                if hh <= 0:
                    continue
                shares.append({"id": h.get("id") or "", "name": h.get("name") or "", "hours": hh})
            assignee = ""
            assignee_id = ""
            if shares:
                primary = max(shares, key=lambda s: (s["hours"], s["name"]))
                assignee = primary.get("name") or ""
                assignee_id = primary.get("id") or ""
            deadline_s = _deadline_iso(deadline, ym)
            prefix = name.split("-", 1)[0].upper()
            excel_status = status_from_line_fill(_cell_fill_hex(ws.cell(r, 1), theme))
            rows.append(
                {
                    "name": name,
                    "line": prefix if prefix in {"DD", "DX", "DS"} else line,
                    "shifts": shifts_n,
                    "hours": hours_n,
                    "deadline": deadline_s,
                    "assigneeName": assignee,
                    "assigneeId": assignee_id,
                    "shares": shares,
                    "excelStatus": excel_status,
                }
            )
        return {"people": people, "rows": rows, "sheet": ws.title}
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _resolved_shares(row: dict, people: list[dict], by_name: dict) -> list[dict]:
    ids = {p["id"] for p in people}
    out = []
    seen: set[str] = set()
    for share in row.get("shares") or row.get("help") or []:
        pid = str(share.get("id") or "").strip()
        if not pid or pid not in ids:
            person = by_name.get(_norm_name(share.get("name") or ""))
            pid = person["id"] if person else ""
        hours = round_hours(share.get("hours"))
        if not pid or pid not in ids or hours <= 0 or pid in seen:
            continue
        seen.add(pid)
        out.append({"id": pid, "hours": hours})
    return out


def _help_from_row(row: dict, people: list[dict], by_name: dict, assignee_id: str) -> list[dict]:
    return [s for s in _resolved_shares(row, people, by_name) if s["id"] != assignee_id]


def _next_task_id(db: dict, prefix: str = "t-xl") -> str:
    used = {str(t.get("id") or "") for t in db.get("tasks") or []}
    n = len(used) + 1
    while f"{prefix}-{n}" in used:
        n += 1
    return f"{prefix}-{n}"


def apply_to_db(db: dict, parsed: dict, ym: str) -> dict:
    """Every PROJECT NAME row on the month tab is upserted. No row is skipped."""
    people = designers(db)
    by_name = people_by_name(people)
    existing = {
        str(t.get("name") or "").strip(): t
        for t in db.get("tasks") or []
        if (t.get("month") or ym) == ym
        and "play" not in str(t.get("department") or t.get("direction") or "").lower()
    }
    by_key = {}
    for task in existing.values():
        key = _id_key(task.get("name") or "")
        if key and key not in by_key:
            by_key[key] = task
    added = 0
    updated = 0
    parsed_names = {str(row.get("name") or "").strip() for row in parsed.get("rows") or []}
    parsed_keys = {_id_key(n) for n in parsed_names if _id_key(n)}
    db["droppedNames"] = [
        n
        for n in (db.get("droppedNames") or [])
        if str(n).strip() not in parsed_names and _id_key(n) not in parsed_keys
    ]
    seen_names: set[str] = set()
    seen_keys: set[str] = set()
    for row in parsed.get("rows") or []:
        name = row["name"]
        key = _id_key(name)
        if name in seen_names or (key and key in seen_keys):
            continue
        seen_names.add(name)
        if key:
            seen_keys.add(key)
        assignee_id = ""
        if row.get("assigneeId") and row["assigneeId"] in {p["id"] for p in people}:
            assignee_id = row["assigneeId"]
        elif row.get("assigneeName"):
            person = by_name.get(_norm_name(row["assigneeName"]))
            if person:
                assignee_id = person["id"]
        shares = _resolved_shares(row, people, by_name)
        task = existing.get(name) or (by_key.get(key) if key else None)
        keep = (task.get("assigneeId") or "") if task else ""
        people_ids = {p["id"] for p in people}
        if keep and keep in people_ids:
            assignee_id = keep
        help_rows = [s for s in shares if s["id"] != assignee_id]
        excel_status = row.get("excelStatus") or ""
        if task:
            old_name = str(task.get("name") or "").strip()
            task["name"] = name
            task["shifts"] = row["shifts"]
            task["hours"] = row["hours"]
            task["line"] = row["line"]
            task["deadline"] = row.get("deadline") or task.get("deadline") or ""
            task["assigneeId"] = assignee_id
            if help_rows:
                task["help"] = help_rows
            task["month"] = ym
            if old_name and old_name != name:
                existing.pop(old_name, None)
            existing[name] = task
            if key:
                by_key[key] = task
            if excel_status:
                task["excelStatus"] = excel_status
            current = task.get("studioStatus") or ""
            if excel_status == "closed":
                task["studioStatus"] = "closed"
                task["mondayStatus"] = "Closed"
            elif excel_status == "wip" and current in {"", "new"}:
                task["studioStatus"] = "wip"
                task["mondayStatus"] = "WIP"
            updated += 1
        else:
            task = {
                "id": _next_task_id(db),
                "name": name,
                "line": row["line"],
                "department": "MGX",
                "game": "",
                "shifts": row["shifts"],
                "hours": row["hours"],
                "deadline": row.get("deadline") or "",
                "month": ym,
                "excelStatus": excel_status,
                "studioStatus": "closed" if excel_status == "closed" else ("wip" if excel_status == "wip" else "new"),
                "mondayStatus": "Closed" if excel_status == "closed" else ("WIP" if excel_status == "wip" else "Ready to Start"),
                "assigneeId": assignee_id,
                "help": help_rows,
                "brief": "",
                "folderUrl": "",
                "docUrl": "",
                "resultUrl": "",
                "source": "excel",
                "activity": [{"at": datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"), "text": "Pulled from SP_MGX_check.xlsx"}],
            }
            db.setdefault("tasks", []).append(task)
            existing[name] = task
            if key:
                by_key[key] = task
            added += 1
    return {
        "added": added,
        "updated": updated,
        "matched": added + updated,
        "parsed": len(parsed.get("rows") or []),
        "people": parsed.get("people") or [],
        "sheet": parsed.get("sheet"),
    }


def tasks_for_finance(db: dict, ym: str) -> list[dict]:
    """Hours and assignees for money from the board. Portal is the source of truth."""
    return [t for t in db.get("tasks") or [] if (t.get("month") or ym) == ym]


def _style_range(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = CALIBRI
        cell.border = THIN
        cell.alignment = CENTER if c != 2 else Alignment(vertical="center")


def _label_section_empl(ws, col: int, n: int) -> None:
    for r in range(2, min((ws.max_row or 2) + 1, 80)):
        a = str(ws.cell(r, 1).value or "").strip().upper()
        if a in {"DX", "DD", "DS"}:
            cell = ws.cell(r, col, f"EMPL {n}")
            cell.fill = HEADER_FILL
            cell.font = CALIBRI_BOLD
            cell.alignment = CENTER
            cell.border = THIN


def _bind_people_columns(ws, people: list[dict]) -> list[tuple[int, dict | None]]:
    """Bind EMPL columns to person id. The header name is only a label and can change."""
    by_id = {p["id"]: p for p in people}
    by_name = people_by_name(people)
    used = set()
    columns: list[tuple[int, dict | None]] = []
    for h in _header_columns(ws):
        col = h["col"]
        cell = ws.cell(1, col)
        person = by_id.get(h["id"]) if h.get("id") else None
        if person and person["id"] in used:
            person = None
        if person is None and h.get("name"):
            candidate = by_name.get(h["name"].strip().lower())
            if candidate and candidate["id"] not in used:
                person = candidate
        if person:
            _stamp_person_header(cell, person)
            used.add(person["id"])
            columns.append((col, person))
        else:
            columns.append((col, None))
    next_col = max((c for c, _ in columns), default=5) + 1
    n = sum(1 for _, p in columns if p) + 1
    for person in people:
        if person["id"] in used:
            continue
        _stamp_person_header(ws.cell(1, next_col), person)
        _label_section_empl(ws, next_col, n)
        ws.column_dimensions[get_column_letter(next_col)].width = 13
        columns.append((next_col, person))
        used.add(person["id"])
        next_col += 1
        n += 1
    return columns


def _write_task_row(ws, row: int, task: dict, columns: list[tuple[int, dict | None]]) -> None:
    shifts = task.get("shifts")
    ws.cell(row, 2, task.get("name") or "")
    ws.cell(row, 3, shifts if shifts not in (None, "") else None)
    ws.cell(row, 4, f"=C{row}*9")
    deadline = task.get("deadline") or ""
    if isinstance(deadline, datetime):
        ws.cell(row, 5, deadline)
        ws.cell(row, 5).number_format = "DD.MM"
    elif deadline:
        try:
            ws.cell(row, 5, datetime.strptime(str(deadline)[:10], "%Y-%m-%d"))
            ws.cell(row, 5).number_format = "DD.MM"
        except ValueError:
            ws.cell(row, 5, deadline)
    else:
        ws.cell(row, 5, None)
    help_rows = help_entries(task)
    last_col = 5
    for col, person in columns:
        last_col = max(last_col, col)
        cell = ws.cell(row, col)
        share = hours_for_person(task, person["id"]) if person else 0.0
        if person and share > 0:
            cell.value = share if help_rows else f"=C{row}*9"
            cell.fill = _hex_fill(person.get("color") or "92D050")
        else:
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
    _style_range(ws, row, last_col)
    mark = line_fill_for_status(task.get("studioStatus") or "")
    if mark is not None:
        ws.cell(row, 1).fill = mark
    ws.cell(row, 2).alignment = Alignment(vertical="center")
    ws.cell(row, 4).alignment = CENTER


def _scan_sheet(ws) -> dict:
    headers = {}
    rows = {}
    total_row = None
    line = "DX"
    last_used = 1
    for r in range(2, (ws.max_row or 1) + 1):
        a = _cell_text(ws.cell(r, 1).value)
        b = _cell_text(ws.cell(r, 2).value)
        if a.lower() == "total":
            if total_row is None:
                total_row = r
            last_used = max(last_used, r)
            continue
        if a.upper() in LINE_MARKS and (not b or b.lower() in ROW_LABELS):
            line = a.upper()
            headers[line] = r
            last_used = r
            continue
        name = _project_name(a, b)
        if not name:
            continue
        rows[_norm_name(name)] = {"row": r, "line": line, "name": name}
        last_used = r
    return {"headers": headers, "rows": rows, "total_row": total_row, "last_used": last_used}


def _insert_in_section(ws, scan: dict, line: str) -> int:
    headers = scan["headers"]
    total_row = scan["total_row"]
    start = headers.get(line)
    later = [r for key, r in headers.items() if start and r > start]
    if later:
        bound = min(later)
    elif total_row:
        bound = total_row
    else:
        bound = (scan["last_used"] or 1) + 1
    last_data = start or 1
    for r in range((start or 1) + 1, bound):
        if _project_name(_cell_text(ws.cell(r, 1).value), _cell_text(ws.cell(r, 2).value)):
            last_data = r
    insert_at = last_data + 1
    ws.insert_rows(insert_at)
    if total_row and total_row >= insert_at:
        scan["total_row"] = total_row + 1
    for key, r in list(headers.items()):
        if r >= insert_at:
            headers[key] = r + 1
    for item in scan["rows"].values():
        if item["row"] >= insert_at:
            item["row"] += 1
    scan["last_used"] = max(scan.get("last_used") or 1, insert_at)
    return insert_at


def _delete_sheet_row(ws, scan: dict, row: int) -> None:
    ws.delete_rows(row)
    if scan.get("total_row") and scan["total_row"] > row:
        scan["total_row"] -= 1
    for key, r in list(scan["headers"].items()):
        if r > row:
            scan["headers"][key] = r - 1
    gone = [key for key, item in scan["rows"].items() if item["row"] == row]
    for item in scan["rows"].values():
        if item["row"] > row:
            item["row"] -= 1
    for key in gone:
        scan["rows"].pop(key, None)
    scan["last_used"] = max((scan.get("last_used") or 1) - 1, 1)


def _drop_dropped_rows(ws, scan: dict, db: dict) -> None:
    dropped = {str(n).strip() for n in (db.get("droppedNames") or []) if n}
    dropped_keys = {_id_key(n) for n in dropped if _id_key(n)}
    if not dropped and not dropped_keys:
        return
    victims = []
    for item in scan["rows"].values():
        name = str(item.get("name") or "").strip()
        if name in dropped or (_id_key(name) and _id_key(name) in dropped_keys):
            victims.append(item["row"])
    for row in sorted(set(victims), reverse=True):
        _delete_sheet_row(ws, scan, row)


def _refresh_total(ws, scan: dict, columns: list[tuple[int, dict | None]]) -> None:
    total_row = scan.get("total_row")
    if not total_row:
        return
    end = total_row - 1
    ws.cell(total_row, 1, "total")
    ws.cell(total_row, 3, f"=SUM(C3:C{end})")
    ws.cell(total_row, 4, f"=SUM(D3:D{end})")
    last_col = 5
    for col, _ in columns:
        last_col = max(last_col, col)
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}3:{letter}{end})")
    _style_range(ws, total_row, last_col)
    for c in range(1, last_col + 1):
        ws.cell(total_row, c).fill = HEADER_FILL
        ws.cell(total_row, c).font = CALIBRI_BOLD


def rewrite_existing_name(db: dict, ym: str, old_name: str, task: dict, path: Path | None = None) -> bool:
    """Change PROJECT NAME on an existing sheet row. Does not add a new row."""
    path = Path(path or XLSX_PATH)
    if not path.exists():
        return False
    wb = load_workbook(path)
    try:
        ws = find_sheet(wb, ym)
        if ws is None:
            return False
        scan = _scan_sheet(ws)
        columns = _bind_people_columns(ws, designers(db))
        found = scan["rows"].get(_norm_name(old_name))
        if not found:
            old_key = _id_key(old_name)
            if old_key:
                found = next(
                    (item for item in scan["rows"].values() if _id_key(item.get("name") or "") == old_key),
                    None,
                )
        if not found:
            return False
        _write_task_row(ws, found["row"], task, columns)
        _atomic_save(wb, path)
        return True
    finally:
        try:
            wb.close()
        except Exception:
            pass


def write_month(db: dict, ym: str, path: Path | None = None) -> Path:
    """Update the month tab in place. Keep formulas (=C*9) and other sheets."""
    path = Path(path or XLSX_PATH)
    if not path.exists():
        raise FileNotFoundError(f"Excel workbook missing: {path}")
    wb = load_workbook(path)
    try:
        return _write_month_body(db, ym, path, wb)
    finally:
        _close_wb(wb)


def _write_month_body(db: dict, ym: str, path: Path, wb) -> Path:
    ws = find_sheet(wb, ym)
    people = designers(db)
    y = ym.split("-")[0]
    if ws is None:
        ws = wb.create_sheet(sheet_name_for(ym), 0)
        ws.cell(1, 2, f"{y} PROJECT NAME")
        for i, person in enumerate(people):
            _stamp_person_header(ws.cell(1, 6 + i), person)
        labels = ["DX", "MGX", "SHIFTS", "HOURS", "DEADLINE"] + [f"EMPL {i + 1}" for i in range(len(people))]
        for c, label in enumerate(labels, 1):
            cell = ws.cell(2, c, label)
            cell.fill = HEADER_FILL
            cell.font = CALIBRI_BOLD
            cell.border = THIN
            cell.alignment = CENTER
        dd_row = 6
        labels[0] = "DD"
        for c, label in enumerate(labels, 1):
            cell = ws.cell(dd_row, c, label)
            cell.fill = HEADER_FILL
            cell.font = CALIBRI_BOLD
            cell.border = THIN
            cell.alignment = CENTER
        ws.cell(10, 1, "total")
        ws.column_dimensions["A"].width = 11
        ws.column_dimensions["B"].width = 63
        for col in range(3, 6 + max(len(people), 1)):
            ws.column_dimensions[get_column_letter(col)].width = 13

    scan = _scan_sheet(ws)
    if "DX" not in scan["headers"]:
        scan["headers"]["DX"] = 2
    columns = _bind_people_columns(ws, people)
    _drop_dropped_rows(ws, scan, db)
    id_index = {
        _id_key(item["name"]): item for item in scan["rows"].values() if _id_key(item["name"])
    }
    tasks = [
        t
        for t in db.get("tasks") or []
        if (t.get("month") or ym) == ym
        and "play" not in str(t.get("department") or t.get("direction") or "").lower()
    ]
    for task in tasks:
        key = _norm_name(task.get("name"))
        kid = _id_key(task.get("name") or "")
        found = scan["rows"].get(key)
        if not found and kid:
            found = id_index.get(kid)
        if found and scan.get("total_row") and found["row"] > scan["total_row"]:
            _delete_sheet_row(ws, scan, found["row"])
            id_index = {
                _id_key(item["name"]): item for item in scan["rows"].values() if _id_key(item["name"])
            }
            found = None
        if found:
            _write_task_row(ws, found["row"], task, columns)
            continue
        line = (task.get("line") or "DD").upper()
        if line not in {"DX", "DD", "DS"}:
            line = "DD"
        if line not in scan["headers"]:
            line = "DD" if "DD" in scan["headers"] else next(iter(scan["headers"]), "DX")
        row = _insert_in_section(ws, scan, line)
        _write_task_row(ws, row, task, columns)
        rec = {"row": row, "line": line, "name": task.get("name")}
        scan["rows"][key] = rec
        if kid:
            id_index[kid] = rec
    _refresh_total(ws, scan, columns)
    _atomic_save(wb, path)
    return path


def _atomic_save(wb, path: Path) -> None:
    """Replace the xlsx only after a full write, so a crash cannot leave a 0-byte sheet."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    tmp_path = Path(tmp)
    try:
        wb.save(tmp_path)
        tmp_path.replace(path)
    except Exception:
        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def _close_wb(wb) -> None:
    try:
        wb.close()
    except Exception:
        pass


def _as_float(value, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def client_hours_filename(ym: str) -> str:
    return f"SUPERPLAY_hours_{ym}.xlsx"


TZ_NAME = re.compile(r"^[A-Za-z0-9_+\-/]{1,64}$")


def parse_tz(name: str) -> ZoneInfo:
    raw = str(name or "").strip()
    if not TZ_NAME.match(raw):
        return ZoneInfo("UTC")
    try:
        return ZoneInfo(raw)
    except Exception:
        return ZoneInfo("UTC")


def zone_city(tz_name: str) -> str:
    part = str(tz_name or "UTC").split("/")[-1].replace("_", " ").strip()
    return part or "UTC"


def local_stamp(now: datetime | None = None, tz_name: str = "UTC") -> str:
    zone = parse_tz(tz_name)
    dt = now or datetime.now(zone)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc).astimezone(zone)
    else:
        dt = dt.astimezone(zone)
    city = zone_city(getattr(zone, "key", None) or tz_name)
    return f"{dt.day} {dt.strftime('%b %Y, %H:%M')} {city}"


def client_hours_bytes(db: dict, ym: str, *, generated_at: datetime | None = None, tz_name: str = "UTC") -> bytes:
    """Build a fresh SuperPlay hours xlsx at download time. No people, rates, or status."""
    ym = str(ym or "").strip()
    shift_hours = _as_float((db.get("meta") or {}).get("shiftHours"), 9) or 9
    tasks = [t for t in db.get("tasks") or [] if (t.get("month") or ym) == ym]
    sections = []
    for dept, prefix in (("MGX", "MGX"), ("Playables", "Playable")):
        bucket = [
            t
            for t in tasks
            if ("play" in str(t.get("department") or t.get("direction") or "").lower())
            == (dept == "Playables")
        ]
        lines = ["DX", "DD"]
        if any(str(t.get("line") or "") == "DS" for t in bucket):
            lines.append("DS")
        for line in lines:
            rows = [t for t in bucket if t.get("line") == line]
            if not rows:
                continue
            sections.append((f"{prefix} {line}", line, rows))
    year = ym[:4] if len(ym) >= 4 else "2026"

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name_for(ym)[:31]
    stamp = ws.cell(1, 1, f"Generated {local_stamp(generated_at, tz_name)}")
    stamp.font = Font(name="Calibri", size=11, color="676879")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    headers = (f"{year} PROJECT NAME", "SHIFTS", "HOURS")
    for col, title in enumerate(headers, 1):
        cell = ws.cell(2, col, title)
        cell.font = CALIBRI_BOLD
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = CENTER

    row = 3

    def write_dir_total(label: str, hours: float) -> None:
        nonlocal row
        ws.cell(row, 1, label)
        ws.cell(row, 3, round(hours, 1))
        for col in range(1, 4):
            cell = ws.cell(row, col)
            cell.font = CALIBRI_TOTAL
            cell.fill = TOTAL_FILL
            cell.border = THIN
        ws.cell(row, 3).alignment = CENTER
        row += 1

    grouped = (("MGX", "MGX TOTAL"), ("Playable", "Playable TOTAL"))
    wrote = 0
    for prefix, total_label in grouped:
        block = [(label, line, rows) for label, line, rows in sections if str(label).startswith(prefix + " ")]
        if not block:
            continue
        if wrote:
            row += 2
        dir_hours = 0.0
        for label, line, rows in block:
            rows = sorted(rows, key=lambda t: str(t.get("name") or "").lower())
            sec_hours = round(sum(client_hours(t, shift_hours) for t in rows), 1)
            dir_hours += sec_hours
            ws.cell(row, 1, label)
            ws.cell(row, 3, sec_hours)
            for col in range(1, 4):
                cell = ws.cell(row, col)
                cell.fill = LINE_FILL.get(line, HEADER_FILL)
                cell.font = CALIBRI_BOLD
                cell.border = THIN
            ws.cell(row, 3).alignment = CENTER
            row += 1
            for task in rows:
                ws.cell(row, 1, str(task.get("name") or ""))
                ws.cell(row, 2, _as_float(task.get("shifts")))
                ws.cell(row, 3, client_hours(task, shift_hours))
                for col in range(1, 4):
                    cell = ws.cell(row, col)
                    cell.font = CALIBRI
                    cell.border = THIN
                ws.cell(row, 2).alignment = CENTER
                ws.cell(row, 3).alignment = CENTER
                row += 1
        write_dir_total(total_label, dir_hours)
        wrote += 1
    ws.column_dimensions["A"].width = 72
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.freeze_panes = "A3"
    buf = BytesIO()
    try:
        wb.save(buf)
    finally:
        _close_wb(wb)
    return buf.getvalue()
