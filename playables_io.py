"""Read/write SP_PLBL_check.xlsx — billed hours, not shifts × 9."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import direction_io
import excel_io

ROOT = Path(__file__).resolve().parent
XLSX_PATH = ROOT / "data" / "SP_PLBL_check.xlsx"
LG_ID = excel_io.LG_ID
CALIBRI = Font(name="Calibri", size=11)
CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(
    left=Side(style="hair", color="B4B4B4"),
    right=Side(style="hair", color="B4B4B4"),
    top=Side(style="hair", color="B4B4B4"),
    bottom=Side(style="hair", color="B4B4B4"),
)
HEADER_FILL = PatternFill("solid", fgColor="D8D8D8")
LINE_MARKS = {"DX", "DD", "DS"}
ROW_LABELS = {"playableads", "plads", "hours", "producer", "project name", "d.line"}


def crew(db: dict) -> list[dict]:
    """Hour columns on SP_PLBL_check.xlsx: AD + DEV.

    MOTION/Сергій stays on the Playables board, but the live sheet has no MOTION
    column. Writing him at col E used to shift Anna's 2DART 1 hours off the row.
    """
    people = [
        p
        for p in db.get("people") or []
        if p.get("role") in {"art_director", "programmer"}
        and direction_io.person_on(p, direction_io.PLAYABLES)
    ]
    return sorted(people, key=lambda p: (p.get("payOrder") or 99, p.get("name") or ""))


def crew_hours(task: dict) -> float:
    """AD + DEV + MOTION seats. No studio buffer."""
    total = 0.0
    total += excel_io.round_hours(task.get("assigneeHours"))
    for row in excel_io.help_entries(task):
        total += row["hours"]
    total += excel_io.round_hours(task.get("programmerHours"))
    for row in task.get("programmerHelp") or []:
        total += excel_io.round_hours((row or {}).get("hours"))
    total += excel_io.round_hours(task.get("motionHours"))
    return round(total, 1)


def producer_hours(task: dict) -> float:
    return crew_hours(task)


def studio_buffer(task: dict) -> float:
    """DL-only extra on top of crew. Excel HOURS − people used to hold this."""
    if task.get("studioBuffer") not in (None, ""):
        return excel_io.round_hours(task.get("studioBuffer"))
    crew = crew_hours(task)
    billed = excel_io.round_hours(task.get("hours"))
    return max(0.0, round(billed - crew, 1))


def billed_hours(task: dict) -> float:
    return round(crew_hours(task) + studio_buffer(task), 1)


def attach_studio_buffer(task: dict) -> bool:
    """Persist Excel leftover as studioBuffer and keep billed = crew + buffer."""
    if not direction_io.is_playables(task):
        return False
    crew = crew_hours(task)
    buf = studio_buffer(task)
    billed = round(crew + buf, 1)
    dirty = False
    if excel_io.round_hours(task.get("studioBuffer")) != buf:
        task["studioBuffer"] = buf
        dirty = True
    if excel_io.round_hours(task.get("hours")) != billed:
        task["hours"] = billed
        dirty = True
    if excel_io.round_hours(task.get("producerHours")) != crew:
        task["producerHours"] = crew
        dirty = True
    return dirty


def attach_studio_buffers(db: dict) -> bool:
    dirty = False
    for task in db.get("tasks") or []:
        if attach_studio_buffer(task):
            dirty = True
    return dirty


def ad_id(task: dict, people: list[dict] | None = None) -> str:
    """Playable AD seat. Programmer is never AD, even if old Excel Who fell back to DEV."""
    aid = str(task.get("assigneeId") or "").strip()
    if not aid:
        return ""
    if aid == str(task.get("programmerId") or "").strip():
        return ""
    if people:
        person = next((p for p in people if str(p.get("id") or "") == aid), None)
        role = (person or {}).get("role")
        if role and role != "art_director":
            return ""
    return aid


def normalize_playable_seats(db: dict) -> bool:
    people = db.get("people") or []
    dirty = False
    for task in db.get("tasks") or []:
        if not direction_io.is_playables(task):
            continue
        nxt = ad_id(task, people)
        if str(task.get("assigneeId") or "") != nxt:
            task["assigneeId"] = nxt
            dirty = True
    return dirty


def hours_for_person(task: dict, person_id: str) -> float:
    pid = str(person_id or "").strip()
    if not pid:
        return 0.0
    if pid == "manager" or pid == str(task.get("producerId") or "").strip():
        return producer_hours(task)
    if pid == str(task.get("programmerId") or "").strip():
        return excel_io.round_hours(task.get("programmerHours"))
    for row in task.get("programmerHelp") or []:
        if str((row or {}).get("id") or "") == pid:
            return excel_io.round_hours((row or {}).get("hours"))
    if pid == str(task.get("motionId") or "").strip():
        return excel_io.round_hours(task.get("motionHours"))
    if pid == ad_id(task):
        return excel_io.round_hours(task.get("assigneeHours"))
    for row in excel_io.help_entries(task):
        if row["id"] == pid:
            return row["hours"]
    return 0.0


def _cell_text(value) -> str:
    if value in (None, ""):
        return ""
    return str(value).strip()


def _num(value) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, str) and value.strip().startswith("="):
        return 0.0
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def _kind_from_label(label: str) -> str:
    raw = str(label or "").strip().lower().replace(" ", "")
    if "prgm" in raw or raw in {"dev", "dev1", "programmer"}:
        return "programmer"
    if "producer" in raw or raw in {"lp", "empl1"}:
        return "producer"
    if "2dart" in raw or "art" in raw or "3dart" in raw:
        return "art"
    if "motion" in raw:
        return "motion"
    return ""


def _header_columns(ws) -> list[dict]:
    cols = []
    for col in range(4, 24):
        name = _cell_text(ws.cell(1, col).value)
        label = _cell_text(ws.cell(2, col).value)
        pid = excel_io._person_id_from_cell(ws.cell(1, col))
        if not name and not label and not pid:
            if cols:
                break
            continue
        cols.append(
            {
                "col": col,
                "name": name,
                "label": label,
                "id": pid,
                "kind": _kind_from_label(label) or _kind_from_label(name),
            }
        )
    return cols


def parse_workbook(path: Path, ym: str) -> dict:
    empty = {"people": [], "rows": [], "sheet": None}
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as e:
        print(f"[playables] parse {path}: {e}", flush=True)
        return empty
    try:
        ws = excel_io.find_sheet(wb, ym)
        if ws is None:
            return empty
        theme = excel_io.theme_rgbs(wb)
        headers = _header_columns(ws)
        people = [h["name"] for h in headers if h["name"]]
        rows = []
        line = "DS"
        for r in range(2, (ws.max_row or 1) + 1):
            a = _cell_text(ws.cell(r, 1).value)
            b = _cell_text(ws.cell(r, 2).value)
            if a.lower() == "total":
                continue
            if a.upper() in LINE_MARKS and (not b or b.lower().replace(" ", "") in ROW_LABELS):
                line = a.upper()
                continue
            name = excel_io._project_name(a, b)
            if not name:
                continue
            hours = _num(ws.cell(r, 3).value)
            art = []
            programmers = []
            motion = None
            for h in headers:
                if h.get("kind") == "producer":
                    continue
                val = _num(ws.cell(r, h["col"]).value)
                if val <= 0:
                    continue
                share = {"id": h.get("id") or "", "name": h.get("name") or "", "hours": val, "kind": h.get("kind") or "art"}
                if share["kind"] == "programmer":
                    programmers.append(share)
                elif share["kind"] == "motion":
                    motion = share
                else:
                    art.append(share)
            lead = art[0] if art else {}
            help_rows = art[1:] if len(art) > 1 else []
            programmer = programmers[0] if programmers else {}
            programmer_help = programmers[1:] if len(programmers) > 1 else []
            prefix = name.split("-", 1)[0].upper()
            rows.append(
                {
                    "name": name,
                    "line": prefix if prefix in LINE_MARKS else line,
                    "hours": hours,
                    "assigneeName": lead.get("name") or "",
                    "assigneeId": lead.get("id") or "",
                    "assigneeHours": lead.get("hours") or 0.0,
                    "help": help_rows,
                    "programmerName": (programmer or {}).get("name") or "",
                    "programmerId": (programmer or {}).get("id") or "",
                    "programmerHours": (programmer or {}).get("hours") or 0.0,
                    "programmerHelp": programmer_help,
                    "motionName": (motion or {}).get("name") or "",
                    "motionId": (motion or {}).get("id") or "",
                    "motionHours": (motion or {}).get("hours") or 0.0,
                    "producerHours": round(
                        sum(s["hours"] for s in art)
                        + ((programmer or {}).get("hours") or 0)
                        + sum(s.get("hours") or 0 for s in programmer_help)
                        + ((motion or {}).get("hours") or 0),
                        1,
                    ),
                    "excelStatus": excel_io.status_from_line_fill(excel_io._cell_fill_hex(ws.cell(r, 1), theme)),
                }
            )
        return {"people": people, "rows": rows, "sheet": ws.title}
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _resolve_id(share: dict, people: list[dict], by_name: dict) -> str:
    ids = {p["id"] for p in people}
    pid = str(share.get("id") or "").strip()
    if pid in ids:
        return pid
    person = by_name.get(excel_io._norm_name(share.get("name") or ""))
    return person["id"] if person else ""


def apply_to_db(db: dict, parsed: dict, ym: str) -> dict:
    people = [p for p in db.get("people") or [] if p.get("role") in direction_io.PLAY_CREW]
    by_name = excel_io.people_by_name(people)
    existing = {
        str(t.get("name") or "").strip(): t
        for t in db.get("tasks") or []
        if (t.get("month") or ym) == ym and direction_io.is_playables(t)
    }
    added = 0
    updated = 0
    for row in parsed.get("rows") or []:
        name = row["name"]
        assignee_id = _resolve_id(
            {"id": row.get("assigneeId"), "name": row.get("assigneeName")},
            people,
            by_name,
        )
        help_rows = []
        for share in row.get("help") or []:
            pid = _resolve_id(share, people, by_name)
            hours = excel_io.round_hours(share.get("hours"))
            if pid and hours and pid != assignee_id:
                help_rows.append({"id": pid, "hours": hours})
        programmer_id = _resolve_id(
            {"id": row.get("programmerId"), "name": row.get("programmerName")},
            people,
            by_name,
        )
        programmer_help = []
        for share in row.get("programmerHelp") or []:
            pid = _resolve_id(share, people, by_name)
            hours = excel_io.round_hours(share.get("hours"))
            if pid and hours and pid != programmer_id:
                programmer_help.append({"id": pid, "hours": hours})
        motion_id = _resolve_id(
            {"id": row.get("motionId"), "name": row.get("motionName")},
            people,
            by_name,
        )
        # AD seat is art hours only. DEV-only rows (Kraken ITR) keep AD empty —
        # billed leftover is studio buffer, not a fake AD on the programmer.
        excel_ad = assignee_id if assignee_id and assignee_id != programmer_id else ""
        excel_status = row.get("excelStatus") or ""
        task = existing.get(name)
        if task:
            before = (task.get("studioStatus"), task.get("excelStatus"), task.get("mondayStatus"))
            task["name"] = name
            task["line"] = row["line"]
            task["department"] = "Playables"
            task["month"] = ym
            # Portal owns seats, billed hours, and studio buffer. Excel does not rewrite them.
            if excel_status:
                task["excelStatus"] = excel_status
            current = direction_io.playable_norm_status(task.get("studioStatus"))
            if excel_status == "closed":
                task["studioStatus"] = "closed"
                task["mondayStatus"] = "Closed"
            elif current not in {"", "new"}:
                # Board pipeline wins. Excel blue/WIP fill is not a reason to bounce WIP → Ready.
                task["studioStatus"] = current
                task["mondayStatus"] = direction_io.PLAY_STATUS.get(current, task.get("mondayStatus"))
            after = (task.get("studioStatus"), task.get("excelStatus"), task.get("mondayStatus"))
            if before != after:
                updated += 1
        else:
            task = {
                "id": excel_io._next_task_id(db, "t-pl"),
                "name": name,
                "line": row["line"],
                "department": "Playables",
                "game": "",
                "shifts": 0,
                "hours": row["hours"],
                "assigneeHours": row.get("assigneeHours") or 0,
                "producerHours": row.get("producerHours") or 0,
                "programmerId": programmer_id,
                "programmerHours": row.get("programmerHours") or 0,
                "programmerHelp": programmer_help,
                "motionId": motion_id,
                "motionHours": row.get("motionHours") or 0,
                "deadline": "",
                "month": ym,
                "excelStatus": excel_status,
                "studioStatus": "closed" if excel_status == "closed" else "new",
                "mondayStatus": "Closed" if excel_status == "closed" else "Ready",
                "assigneeId": excel_ad,
                "help": help_rows,
                "brief": "",
                "folderUrl": "",
                "docUrl": "",
                "resultUrl": "",
                "source": "excel",
                "activity": [
                    {
                        "at": datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
                        "text": "Pulled from SP_PLBL_check.xlsx",
                    }
                ],
            }
            db.setdefault("tasks", []).append(task)
            attach_studio_buffer(task)
            existing[name] = task
            added += 1
    if normalize_playable_seats(db):
        updated += 1
    if attach_studio_buffers(db):
        updated += 1
    return {
        "added": added,
        "updated": updated,
        "matched": added + updated,
        "parsed": len(parsed.get("rows") or []),
        "people": parsed.get("people") or [],
        "sheet": parsed.get("sheet"),
    }


def _stamp(cell, person: dict) -> None:
    cell.value = person.get("name") or cell.value
    cell.font = CALIBRI_BOLD
    cell.alignment = CENTER
    cell.comment = Comment(f"lg-id:{person['id']}", "LG")


def _drop_inserted_motion(ws, people: list[dict]) -> None:
    """Remove a MOTION column we inserted in front of 2DART/PRGMM. Keep a trailing MOTION seat."""
    crew_ids = {p["id"] for p in people}
    by_name = excel_io.people_by_name(people)
    headers = _header_columns(ws)
    drop: list[int] = []
    for h in headers:
        if h.get("kind") != "motion":
            continue
        pid = str(h.get("id") or "").strip()
        if pid in crew_ids or by_name.get(excel_io._norm_name(h.get("name") or "")):
            continue
        right = [
            x
            for x in headers
            if x["col"] > h["col"]
            and (
                x.get("kind") in {"art", "programmer"}
                or by_name.get(excel_io._norm_name(x.get("name") or ""))
            )
        ]
        if right:
            drop.append(int(h["col"]))
    for col in sorted(set(drop), reverse=True):
        ws.delete_cols(col)


def _bind_crew_columns(ws, people: list[dict]) -> list[tuple[int, dict]]:
    """Match people to existing sheet columns by lg-id or name. Never insert MOTION in the middle."""
    headers = [h for h in _header_columns(ws) if h.get("kind") != "producer"]
    by_id = {p["id"]: p for p in people}
    by_name = excel_io.people_by_name(people)
    bound: list[tuple[int, dict]] = []
    used_ids: set[str] = set()
    used_cols: set[int] = set()
    for h in headers:
        person = None
        pid = str(h.get("id") or "").strip()
        if pid in by_id:
            person = by_id[pid]
        else:
            person = by_name.get(excel_io._norm_name(h.get("name") or ""))
        if not person or person["id"] in used_ids:
            continue
        col = int(h["col"])
        bound.append((col, person))
        used_ids.add(person["id"])
        used_cols.add(col)
    last = max([c for c, _ in bound] + [h["col"] for h in _header_columns(ws)] + [4])
    for person in people:
        if person["id"] in used_ids:
            continue
        last += 1
        while last in used_cols:
            last += 1
        bound.append((last, person))
        used_ids.add(person["id"])
        used_cols.add(last)
    bound.sort(key=lambda pair: pair[0])
    return bound


def _role_label(person: dict, used: dict) -> str:
    role = person.get("role")
    if role == "programmer":
        n = used.get("dev", 0) + 1
        used["dev"] = n
        return f"PRGMM{n}"
    if role == "designer":
        return "MOTION"
    n = used.get("art", 0) + 1
    used["art"] = n
    return f"2DART {n}"


def write_month(db: dict, ym: str, path: Path | None = None) -> Path | None:
    path = Path(path or XLSX_PATH)
    if not path.exists():
        return None
    people = crew(db)
    tasks = direction_io.month_tasks(db, ym, direction_io.PLAYABLES)
    wb = load_workbook(path)
    try:
        ws = excel_io.find_sheet(wb, ym)
        y = ym.split("-")[0]
        if ws is None:
            ws = wb.create_sheet(excel_io.sheet_name_for(ym), 0)
        ws.cell(1, 2, f"{y} PROJECT NAME")
        ws.cell(1, 2).font = CALIBRI_BOLD
        _drop_inserted_motion(ws, people)
        columns = _bind_crew_columns(ws, people)
        used: dict = {}
        for col, person in columns:
            _stamp(ws.cell(1, col), person)
            lab = ws.cell(2, col)
            if not _cell_text(lab.value):
                lab.value = _role_label(person, used)
            lab.fill = HEADER_FILL
            lab.font = CALIBRI_BOLD
            lab.alignment = CENTER
            lab.border = THIN
        prod = ws.cell(
            1,
            4,
            next((p.get("name") for p in db.get("people") or [] if p.get("role") == "manager"), "Anastasia"),
        )
        prod.font = CALIBRI_BOLD
        prod.alignment = CENTER
        lab = ws.cell(2, 4)
        if not _cell_text(lab.value):
            lab.value = "Producer"
        lab.fill = HEADER_FILL
        lab.font = CALIBRI_BOLD
        lab.alignment = CENTER
        ws.cell(2, 1, "DX")
        ws.cell(2, 2, "PlayableADs")
        ws.cell(2, 3, "HOURS")
        last_header = max([c for c, _ in columns] + [4])
        for c in range(1, last_header + 1):
            cell = ws.cell(2, c)
            cell.fill = HEADER_FILL
            cell.font = CALIBRI_BOLD
            cell.border = THIN
            cell.alignment = CENTER

        scan = _scan_sheet(ws)
        excel_io._drop_dropped_rows(ws, scan, db)
        id_index = {
            excel_io._id_key(item["name"]): item
            for item in scan["rows"].values()
            if excel_io._id_key(item.get("name") or "")
        }
        for task in tasks:
            key = excel_io._norm_name(task.get("name") or "")
            kid = excel_io._id_key(task.get("name") or "")
            found = scan["rows"].get(key)
            if not found and kid:
                found = id_index.get(kid)
            if found and scan.get("total_row") and found["row"] > scan["total_row"]:
                excel_io._delete_sheet_row(ws, scan, found["row"])
                id_index = {
                    excel_io._id_key(item["name"]): item
                    for item in scan["rows"].values()
                    if excel_io._id_key(item.get("name") or "")
                }
                found = None
            if found:
                _write_task_row(ws, found["row"], task, columns)
                continue
            line = (task.get("line") or "DD").upper()
            if line not in {"DX", "DD", "DS"}:
                line = "DD"
            if line not in scan["headers"]:
                _ensure_line_header(ws, scan, line, last_header)
            if line not in scan["headers"]:
                line = "DD" if "DD" in scan["headers"] else next(iter(scan["headers"]), "DX")
            row = excel_io._insert_in_section(ws, scan, line)
            _write_task_row(ws, row, task, columns)
            rec = {"row": row, "line": line, "name": task.get("name")}
            scan["rows"][key] = rec
            if kid:
                id_index[kid] = rec
        _refresh_total(ws, scan, columns)
        excel_io._atomic_save(wb, path)
        return path
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _scan_sheet(ws) -> dict:
    headers = {}
    rows = {}
    total_row = None
    line = "DX"
    last_used = 1
    for r in range(2, (ws.max_row or 1) + 1):
        a = _cell_text(ws.cell(r, 1).value)
        b = _cell_text(ws.cell(r, 2).value)
        if a.lower() == "total":
            if total_row is None:
                total_row = r
            last_used = max(last_used, r)
            continue
        if a.upper() in LINE_MARKS and (not b or b.lower().replace(" ", "") in ROW_LABELS):
            line = a.upper()
            headers[line] = r
            last_used = r
            continue
        name = excel_io._project_name(a, b)
        if not name:
            continue
        rows[excel_io._norm_name(name)] = {"row": r, "line": line, "name": name}
        last_used = r
    return {"headers": headers, "rows": rows, "total_row": total_row, "last_used": last_used}


def _ensure_line_header(ws, scan: dict, line: str, last_header: int) -> None:
    if line in scan["headers"]:
        return
    insert_at = scan.get("total_row") or ((scan.get("last_used") or 1) + 1)
    ws.insert_rows(insert_at)
    if scan.get("total_row") and scan["total_row"] >= insert_at:
        scan["total_row"] = scan["total_row"] + 1
    for key, r in list(scan["headers"].items()):
        if r >= insert_at:
            scan["headers"][key] = r + 1
    for item in scan["rows"].values():
        if item["row"] >= insert_at:
            item["row"] += 1
    ws.cell(insert_at, 1, line)
    ws.cell(insert_at, 2, "PlayableADs")
    ws.cell(insert_at, 3, "HOURS")
    for c in range(1, max(last_header, 4) + 1):
        cell = ws.cell(insert_at, c)
        cell.fill = HEADER_FILL
        cell.font = CALIBRI_BOLD
        cell.border = THIN
        cell.alignment = CENTER
    scan["headers"][line] = insert_at
    scan["last_used"] = max(scan.get("last_used") or 1, insert_at)


def _refresh_total(ws, scan: dict, columns: list[tuple[int, dict]]) -> None:
    total_row = scan.get("total_row")
    if not total_row:
        return
    end = max(total_row - 1, 3)
    hour_cols = [col for col, _ in columns]
    last = max(hour_cols + [4, 3])
    ws.cell(total_row, 1, "total")
    ws.cell(total_row, 3, f"=SUM(C3:C{end})")
    ws.cell(total_row, 4, f"=SUM(D3:D{end})")
    for col in hour_cols:
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}3:{letter}{end})")
    for c in range(1, last + 1):
        cell = ws.cell(total_row, c)
        cell.fill = HEADER_FILL
        cell.font = CALIBRI_BOLD
        cell.border = THIN
        cell.alignment = CENTER


def rewrite_existing_name(db: dict, ym: str, old_name: str, task: dict, path: Path | None = None) -> bool:
    """Change PROJECT NAME on an existing Playables sheet row. Does not add a new row."""
    path = Path(path or XLSX_PATH)
    if not path.exists():
        return False
    wb = load_workbook(path)
    try:
        ws = excel_io.find_sheet(wb, ym)
        if ws is None:
            return False
        scan = _scan_sheet(ws)
        columns = _bind_crew_columns(ws, crew(db))
        found = scan["rows"].get(excel_io._norm_name(old_name))
        if not found:
            old_key = excel_io._id_key(old_name)
            if old_key:
                found = next(
                    (item for item in scan["rows"].values() if excel_io._id_key(item.get("name") or "") == old_key),
                    None,
                )
        if not found:
            return False
        _write_task_row(ws, found["row"], task, columns)
        excel_io._atomic_save(wb, path)
        return True
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _write_task_row(ws, row: int, task: dict, columns: list[tuple[int, dict]]) -> None:
    hours = billed_hours(task)
    ws.cell(row, 2, task.get("name") or "")
    ws.cell(row, 3, hours if hours else None)
    hour_cols = [col for col, _ in columns]
    if hour_cols:
        start = get_column_letter(min(hour_cols))
        end = get_column_letter(max(hour_cols))
        ws.cell(row, 4, f"=SUM({start}{row}:{end}{row})")
    last = max(hour_cols + [4])
    for col, person in columns:
        last = max(last, col)
        cell = ws.cell(row, col)
        share = hours_for_person(task, person["id"])
        cell.value = share if share else None
        if share:
            cell.fill = excel_io._hex_fill(person.get("color") or "92D050")
        else:
            cell.fill = PatternFill(fill_type=None)
        cell.alignment = CENTER
        cell.border = THIN
        cell.font = CALIBRI
    for c in range(1, last + 1):
        ws.cell(row, c).border = THIN
        ws.cell(row, c).font = CALIBRI
    ws.cell(row, 2).alignment = Alignment(vertical="center")
    ws.cell(row, 3).alignment = CENTER
    mark = excel_io.line_fill_for_status(task.get("studioStatus") or "")
    if mark is not None:
        ws.cell(row, 1).fill = mark
