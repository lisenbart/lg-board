#!/usr/bin/env python3
"""Crash tests for the studio desk. Run: python3 tools/test_portal.py"""
from __future__ import annotations

import json
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
import urllib.error
import urllib.request
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook

import auth_io
import direction_io
import drive_io
import excel_io
import finance_io
import playables_io
import server

PASS = 0
FAIL = 0
ERRORS: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  ok  {name}")
    else:
        FAIL += 1
        msg = f"  FAIL  {name}" + (f" — {detail}" if detail else "")
        print(msg)
        ERRORS.append(msg)


def _synth_sheet(path: Path, extra_below_total: bool = True, duplicate: bool = False) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "October 26"
    ws.cell(1, 2, "2026 PROJECT NAME")
    ws.cell(2, 1, "DX")
    ws.cell(2, 2, "MGX")
    ws.cell(2, 3, "SHIFTS")
    ws.cell(2, 4, "HOURS")
    ws.cell(2, 5, "DEADLINE")
    ws.cell(3, 2, "DX-AAA-0001")
    ws.cell(3, 3, 1)
    ws.cell(4, 2, "DX-BBB-0002")
    ws.cell(4, 3, 0.5)
    ws.cell(6, 1, "DD")
    ws.cell(6, 2, "MGX")
    ws.cell(6, 3, "SHIFTS")
    ws.cell(6, 4, "HOURS")
    ws.cell(6, 5, "DEADLINE")
    ws.cell(7, 2, "DD-CCC-0003")
    ws.cell(7, 3, 2)
    ws.cell(8, 2, "DD-DDD-0004")
    ws.cell(8, 3, 0)
    ws.cell(10, 1, "total")
    ws.cell(10, 3, "=SUM(C3:C9)")
    if extra_below_total:
        ws.cell(12, 2, "DD-TEST-0000001")
        ws.cell(12, 3, 0)
        ws.cell(13, 1, "DD-TEST-0000002")
        ws.cell(13, 3, 0)
    if duplicate:
        ws.cell(14, 2, "DD-TEST-0000001")
        ws.cell(14, 3, 9)
    wb.save(path)
    wb.close()


def _db() -> dict:
    return {
        "meta": {"month": "2026-10", "rev": 1},
        "people": [
            {"id": "designer", "name": "Masha", "role": "designer", "color": "92D050"},
        ],
        "tasks": [],
        "droppedNames": ["DD-TEST-0000001"],
        "events": [],
    }


def test_client_hours_xlsx() -> None:
    db = {
        "meta": {"shiftHours": 9, "month": "2026-10", "clientRateUsd": 35, "playableClientRateUsd": 40},
        "tasks": [
            {"name": "DX-AAA-0001", "line": "DX", "month": "2026-10", "shifts": 0.3, "hours": 2.7},
            {"name": "DD-BBB-0002", "line": "DD", "month": "2026-10", "shifts": 1, "hours": 9},
            {"name": "DX-OLD", "line": "DX", "month": "2026-09", "shifts": 9, "hours": 81},
            {
                "name": "DS-P-060_MergeSolitaire",
                "line": "DS",
                "month": "2026-10",
                "department": "Playables",
                "shifts": 0,
                "hours": 30,
            },
            {
                "name": "DS-P-039_Kraken_ITR",
                "line": "DS",
                "month": "2026-10",
                "shifts": 0,
                "hours": 5,
            },
        ],
    }
    at = datetime(2026, 10, 14, 18, 29, tzinfo=ZoneInfo("Europe/Kyiv"))
    data = excel_io.client_hours_bytes(db, "2026-10", generated_at=at, tz_name="Europe/Kyiv")
    check("hours_direction uses name when dept missing", excel_io.hours_direction({"name": "DS-P-1", "line": "DS"}) == "playables")
    check("hours_direction keeps MGX without P-", excel_io.hours_direction({"name": "DS-AIV-1", "line": "DS"}) == "mgx")
    check("hours workbook is xlsx", data[:2] == b"PK")
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    col_a = [str(ws.cell(r, 1).value or "") for r in range(1, (ws.max_row or 1) + 1)]
    check("hours stamp uses viewer zone", "Generated 14 Oct 2026, 18:29 Kyiv" in col_a[0], col_a[0])
    tel = excel_io.local_stamp(at, "Asia/Tel_Aviv")
    check("hours stamp city follows tz", "Tel Aviv" in tel, tel)
    check("hours sheet skips other months", "DX-OLD" not in col_a)
    check("hours sheet has both lines", "DX-AAA-0001" in col_a and "DD-BBB-0002" in col_a)
    mgx_total = next(
        ws.cell(r, 3).value
        for r in range(1, (ws.max_row or 1) + 1)
        if str(ws.cell(r, 1).value or "").startswith("MGX TOTAL")
    )
    play_total = next(
        ws.cell(r, 3).value
        for r in range(1, (ws.max_row or 1) + 1)
        if str(ws.cell(r, 1).value or "").startswith("Playable TOTAL")
    )
    check("hours MGX TOTAL is 11.7", float(mgx_total) == 11.7, str(mgx_total))
    check("hours Playable TOTAL is 35", float(play_total) == 35, str(play_total))
    check("hours playable DS is not MGX DS", "MGX DS" not in col_a, str([n for n in col_a if " DS" in n or n.startswith("DS-P-")]))
    check("hours playable DS uses Playable label", "Playable DS" in col_a, str(col_a))
    check("hours nameless-dept playable stays Playable", "DS-P-039_Kraken_ITR" in col_a)
    check(
        "hours totals omit rates",
        all("$" not in n and "/h" not in n for n in col_a if "TOTAL" in n),
        str([n for n in col_a if "TOTAL" in n]),
    )
    mgx_row = next(r for r in range(1, (ws.max_row or 1) + 1) if str(ws.cell(r, 1).value or "") == "MGX TOTAL")
    play_row = next(r for r in range(1, (ws.max_row or 1) + 1) if str(ws.cell(r, 1).value or "") == "Playable TOTAL")
    mgx_fill = str(getattr(ws.cell(mgx_row, 3).fill.fgColor, "rgb", "") or "")
    sec_row = next(r for r in range(1, (ws.max_row or 1) + 1) if str(ws.cell(r, 1).value or "") == "Playable DS")
    sec_fill = str(getattr(ws.cell(sec_row, 3).fill.fgColor, "rgb", "") or "")
    check("hours MGX TOTAL uses darkest pastel", mgx_fill.upper().endswith("9BB5CE"), mgx_fill)
    check("hours section sum is mid pastel", sec_fill.upper().endswith("CAD9E8"), sec_fill)
    check("hours fills are not WIP blue", "579BFC" not in mgx_fill.upper() and "579BFC" not in sec_fill.upper())
    check("hours blocks are separate", play_row - mgx_row >= 3, f"{mgx_row} → {play_row}")
    check("hours sheet has no mixed TOTAL", "TOTAL" not in col_a)
    empl = any("EMPL" in str(ws.cell(2, c).value or "") for c in range(1, 8))
    check("hours sheet has no EMPL", not empl)
    wb.close()
    later = excel_io.client_hours_bytes(
        db, "2026-10", generated_at=datetime(2026, 10, 14, 19, 5, tzinfo=ZoneInfo("Europe/Kyiv")), tz_name="Europe/Kyiv"
    )
    a = load_workbook(BytesIO(data)).active.cell(1, 1).value
    b = load_workbook(BytesIO(later)).active.cell(1, 1).value
    check("hours xlsx is generated per download", a != b, f"{a} vs {b}")


def test_pin_opens() -> None:
    check(
        "DL PIN opens every studio seat",
        set(auth_io.pin_owner_roles("3245"))
        == {"finance", "manager", "teamlead", "designer", "art_director", "programmer"},
    )
    check(
        "LP PIN opens LP and below",
        set(auth_io.pin_owner_roles("1726"))
        == {"manager", "teamlead", "designer", "art_director", "programmer"},
    )
    check(
        "TL PIN opens TL and performers",
        set(auth_io.pin_owner_roles("2026")) == {"teamlead", "designer", "art_director", "programmer"},
    )
    check(
        "team PIN opens MD AD DEV",
        set(auth_io.pin_owner_roles("4429")) == {"designer", "art_director", "programmer"},
    )
    check("bad PIN opens nothing", auth_io.pin_owner_roles("0000") == ())
    people = [
        {"id": "dl", "name": "D", "role": "finance"},
        {"id": "lp", "name": "A", "role": "manager"},
        {"id": "tl", "name": "N", "role": "teamlead"},
        {"id": "md", "name": "M", "role": "designer", "payOrder": 2},
        {"id": "md2", "name": "S", "role": "designer", "payOrder": 3},
        {"id": "ad", "name": "Y", "role": "art_director"},
    ]
    dl_md = auth_io.resolve_studio_login(people, "3245", "designer")
    check(
        "DL PIN on MD door asks which designer",
        dl_md.get("needSeat") is True and [s.get("id") for s in dl_md.get("seats") or []] == ["md", "md2"],
        str(dl_md),
    )
    dl_pick = auth_io.resolve_studio_login(people, "3245", "designer", "md2")
    check("DL PIN sits the MD you picked", (dl_pick.get("person") or {}).get("id") == "md2", str(dl_pick))
    dl_lp = auth_io.resolve_studio_login(people, "3245", "manager")
    check("DL PIN on LP door sits LP", (dl_lp.get("person") or {}).get("id") == "lp", str(dl_lp))
    lp_dl = auth_io.resolve_studio_login(people, "1726", "finance")
    check("LP PIN cannot open DL", "cannot open" in str(lp_dl.get("error") or ""), str(lp_dl))
    lp_md = auth_io.resolve_studio_login(people, "1726", "designer")
    check(
        "LP PIN on MD door asks which designer",
        lp_md.get("needSeat") is True and [s.get("id") for s in lp_md.get("seats") or []] == ["md", "md2"],
        str(lp_md),
    )
    lp_tl = auth_io.resolve_studio_login(people, "1726", "teamlead")
    check("LP PIN on TL door sits TL", (lp_tl.get("person") or {}).get("id") == "tl", str(lp_tl))
    tl_lp = auth_io.resolve_studio_login(people, "2026", "manager")
    check("TL PIN cannot open LP", "cannot open" in str(tl_lp.get("error") or ""), str(tl_lp))
    tl_ad = auth_io.resolve_studio_login(people, "2026", "art_director")
    check("TL PIN on AD door sits AD", (tl_ad.get("person") or {}).get("id") == "ad", str(tl_ad))
    team_md = auth_io.resolve_studio_login(people, "4429", "designer")
    team_ids = [s.get("id") for s in (team_md.get("seats") or [])]
    check(
        "team PIN on MD asks which designer",
        team_md.get("needSeat") is True and team_ids == ["md", "md2"],
        str(team_md),
    )
    team_ad = auth_io.resolve_studio_login(people, "4429", "art_director")
    check("team PIN on AD sits the only AD", (team_ad.get("person") or {}).get("id") == "ad", str(team_ad))
    team_dl = auth_io.resolve_studio_login(people, "4429", "finance")
    check("team PIN cannot open DL", "cannot open" in str(team_dl.get("error") or ""), str(team_dl))


def test_person_pins() -> None:
    with tempfile.TemporaryDirectory() as raw:
        orig = auth_io.PINS_PATH
        auth_io.PINS_PATH = Path(raw) / "pins.json"
        try:
            people = [
                {"id": "dl", "name": "D", "role": "finance"},
                {"id": "ad", "name": "Anna", "role": "art_director"},
                {"id": "ad2", "name": "Yulia", "role": "art_director"},
            ]
            check("role PIN fills the Team row", auth_io.effective_pin(people[1]) == "4429")
            err = auth_io.set_person_pin(people, "ad", "8811")
            check("personal PIN saves", err == "", err)
            check("Anna PIN is personal", auth_io.effective_pin(people[1]) == "8811")
            check("Yulia still has the team PIN", auth_io.effective_pin(people[2]) == "4429")
            check(
                "letters are not a PIN",
                auth_io.set_person_pin(people, "ad", "12ab") == "PIN must be 4–8 digits",
            )
            sit = auth_io.resolve_studio_login(people, "8811", "art_director")
            check("unique PIN sits that AD", (sit.get("person") or {}).get("id") == "ad", str(sit))
            blocked = auth_io.resolve_studio_login(people, "8811", "finance")
            check("AD PIN cannot open DL", "cannot open" in str(blocked.get("error") or ""), str(blocked))
            mgx = auth_io.resolve_studio_login(people, "8811", "designer")
            check("personal AD PIN cannot open MD", "cannot open" in str(mgx.get("error") or ""), str(mgx))
            team = auth_io.resolve_studio_login(people, "4429", "art_director")
            check(
                "old team PIN sits only the AD who did not change",
                (team.get("person") or {}).get("id") == "ad2" and not team.get("needSeat"),
                str(team),
            )
            old_anna = auth_io.resolve_studio_login(people, "4429", "art_director", "ad")
            check(
                "old PIN cannot sit a person who changed it",
                "wrong" in str(old_anna.get("error") or "").lower() or (old_anna.get("person") or {}).get("id") != "ad",
                str(old_anna),
            )
            auth_io.set_person_pin(people, "ad", "4430")
            auth_io.set_person_pin(people, "ad2", "4430")
            gone = auth_io.resolve_studio_login(people, "4429", "art_director")
            gone_err = str(gone.get("error") or "")
            check(
                "old team PIN is dead once every AD changed",
                "cannot open" in gone_err or "no seat" in gone_err or "wrong" in gone_err,
                str(gone),
            )
            play = auth_io.resolve_studio_login(people, "4430", "art_director")
            play_ids = [s.get("id") for s in (play.get("seats") or [])]
            if play.get("person"):
                play_ids = [play["person"].get("id")]
            check(
                "new shared PIN sits the Playables ADs",
                set(play_ids) == {"ad", "ad2"},
                str(play),
            )
            auth_io.set_person_pin(people, "ad", "8811")
            auth_io.drop_person_pin("ad2")
            auth_io.set_person_pin(people, "dl", "9999")
            check("solo desk PIN updates the door", auth_io.pin_ok_for_role("finance", "9999"))
            db = {"meta": {}, "people": people, "tasks": []}
            fin = auth_io.view_state(db, {"role": "finance", "email": "dmytro@lisenbart.games"})
            pins = {p.get("id"): p.get("pin") for p in fin.get("people") or []}
            check("finance Team sees PINs", pins.get("ad") == "8811" and pins.get("ad2") == "4429", str(pins))
            lp = auth_io.view_state(db, {"role": "art_director", "email": "lisenbart.games@gmail.com"})
            check("PIN is not on performer state", "pin" not in ((lp.get("people") or [{}])[0]), str(lp.get("people")))
        finally:
            auth_io.PINS_PATH = orig
    js = (ROOT / "public" / "assets" / "app.js").read_text(encoding="utf-8")
    css = (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8")
    check("Team row has a PIN box", 'class="team-pin"' in js and 'placeholder="PIN"' in js)
    check("PIN saves with the name", "payload.pin" in js and "PIN must be 4–8 digits" in js)
    check("PIN box sits next to Name", "7.5rem" in css and ".team-pin" in css)
    check("Team Save waits for edits", "refreshTeamSaves" in js and "teamRowDirty" in js)
    check("card Save waits for edits", "refreshCardSave" in js and "cardFingerprint" in js)
    check("card Save turns ready when dirty", 'classList.toggle("is-ready"' in js)
    check("idle commit buttons start off", "bindLoginGo" in js and "refreshCreateGo" in js)


def test_parse_past_total() -> None:
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "sheet.xlsx"
        _synth_sheet(path)
        parsed = excel_io.parse_workbook(path, "2026-10")
        names = [r["name"] for r in parsed["rows"]]
        check("parse finds sheet", parsed.get("sheet") == "October 26")
        check("parse keeps rows above total", "DX-AAA-0001" in names and "DD-CCC-0003" in names)
        check("parse keeps rows below total", "DD-TEST-0000001" in names and "DD-TEST-0000002" in names)
        check("parse all six projects", len(names) == 6, str(names))
        test_a = next(r for r in parsed["rows"] if r["name"] == "DD-TEST-0000002")
        check("name in column A still counts", test_a["line"] == "DD")


def test_apply_full_sheet() -> None:
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "sheet.xlsx"
        _synth_sheet(path, duplicate=True)
        parsed = excel_io.parse_workbook(path, "2026-10")
        db = _db()
        stats = excel_io.apply_to_db(db, parsed, "2026-10")
        names = [t["name"] for t in db["tasks"]]
        check("apply upserts every distinct row", stats["added"] == 6, str(stats))
        check("duplicate Excel row does not double-card", names.count("DD-TEST-0000001") == 1)
        check("Excel presence un-drops the name", "DD-TEST-0000001" not in db["droppedNames"])
        ids = [t["id"] for t in db["tasks"]]
        check("new cards get unique ids", len(ids) == len(set(ids)), str(ids))
        db["tasks"] = [t for t in db["tasks"] if t["name"] != "DX-BBB-0002"]
        extra = excel_io.apply_to_db(db, parsed, "2026-10")
        ids2 = [t["id"] for t in db["tasks"]]
        check("re-add after delete still unique ids", extra["added"] == 1 and len(ids2) == len(set(ids2)))


def test_write_month_roundtrip() -> None:
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "sheet.xlsx"
        _synth_sheet(path)
        parsed = excel_io.parse_workbook(path, "2026-10")
        db = _db()
        excel_io.apply_to_db(db, parsed, "2026-10")
        excel_io.write_month(db, "2026-10", path)
        again = excel_io.parse_workbook(path, "2026-10")
        names = [r["name"] for r in again["rows"]]
        check("roundtrip keeps TEST rows", "DD-TEST-0000001" in names and "DD-TEST-0000002" in names)
        wb = load_workbook(path)
        ws = wb[again["sheet"]]
        total_row = next(r for r in range(2, (ws.max_row or 1) + 1) if str(ws.cell(r, 1).value or "").strip().lower() == "total")
        test_rows = []
        for r in range(2, (ws.max_row or 1) + 1):
            name = excel_io._project_name(str(ws.cell(r, 1).value or ""), str(ws.cell(r, 2).value or ""))
            if "TEST" in name:
                test_rows.append(r)
        wb.close()
        check("write_month moves TEST above total", test_rows and max(test_rows) < total_row, f"rows={test_rows} total={total_row}")
        check("xlsx still opens after atomic save", path.stat().st_size > 1000)


def test_corrupt_xlsx_empty() -> None:
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "bad.xlsx"
        path.write_text("not a workbook", encoding="utf-8")
        parsed = excel_io.parse_workbook(path, "2026-10")
        check("corrupt xlsx returns empty, does not raise", parsed.get("rows") == [] and not parsed.get("sheet"))


def test_live_workbook_if_present() -> None:
    live = ROOT / "data" / "SP_MGX_check.xlsx"
    if not live.exists():
        check("live Excel present", False, "missing data/SP_MGX_check.xlsx")
        return
    parsed = excel_io.parse_workbook(live, "2026-10")
    names = [r["name"] for r in parsed.get("rows") or []]
    check("live October tab parses", bool(parsed.get("sheet")), str(parsed.get("sheet")))
    check("live October has project rows", len(names) >= 20, str(len(names)))
    db = json.loads((ROOT / "data" / "seed.json").read_text(encoding="utf-8"))
    db.setdefault("droppedNames", [])
    stats = excel_io.apply_to_db(db, parsed, "2026-10")
    board = {
        str(t.get("name") or "").strip()
        for t in db.get("tasks") or []
        if (t.get("month") or "2026-10") == "2026-10"
    }
    missing = [n for n in names if n not in board]
    check("every live Excel row has a card", not missing, str(missing[:8]))
    check("matched count equals parsed names", stats["matched"] == len(set(names)), str(stats))


def _empl_n_headers(path: Path) -> list[str]:
    wb = load_workbook(path)
    found = []
    try:
        for ws in wb.worksheets:
            max_r = min(ws.max_row or 1, 80)
            max_c = min(ws.max_column or 1, 20)
            for r in range(1, max_r + 1):
                b = excel_io._cell_text(ws.cell(r, 2).value)
                if excel_io.PROJECT_NAME_RE.match(b or ""):
                    continue
                for c in range(1, max_c + 1):
                    cell = ws.cell(r, c)
                    if excel_io.is_empl_seat(cell.value):
                        found.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    finally:
        wb.close()
    return found


def test_mgx_seat_headers() -> None:
    db = json.loads((ROOT / "data" / "db.json").read_text(encoding="utf-8"))
    live = ROOT / "data" / "SP_MGX_check.xlsx"
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "mgx.xlsx"
        shutil.copy(live, path)
        excel_io.write_month(db, "2026-10", path)
        wb = load_workbook(path)
        ws = excel_io.find_sheet(wb, "2026-10")
        labels = [str(ws.cell(2, c).value or "") for c in range(6, 12)]
        joined = " ".join(labels)
        dec = next((s for s in wb.sheetnames if "December 25" in s), None)
        dec_e = str(wb[dec].cell(2, 5).value or "") if dec else ""
        wb.close()
        check(
            "MGX check seats are TL / MD not EMPL",
            "TL 1" in labels and "MD 1" in labels and "EMPL" not in joined,
            str(labels),
        )
        leftover = _empl_n_headers(path)
        check("MGX check has no leftover EMPL N seats", not leftover, str(leftover[:12]))
        if dec:
            check("December 25 seat row has no EMPL 1", not excel_io.is_empl_seat(dec_e), dec_e)
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "calc.xlsx"
        shutil.copy(finance_io.CALC_PATH, path)
        finance_io.write_month(db, "2026-10", path=path)
        wb = load_workbook(path)
        ws = wb[finance_io.sheet_name_for("2026-10")]
        labels = [str(ws.cell(5, c).value or "") for c in range(6, 13)]
        wb.close()
        check(
            "MGX calculator seats are LP TL MD",
            labels[:4] == ["LP 1", "TL 1", "MD 1", "MD 2"] and "EMPL 1" not in labels,
            str(labels),
        )
        leftover = _empl_n_headers(path)
        check("MGX calculator has no leftover EMPL N seats", not leftover, str(leftover[:12]))


def test_corrupt_db_recovery() -> None:
    with tempfile.TemporaryDirectory() as raw:
        data = Path(raw)
        db_path = data / "db.json"
        bak = data / "db.json.bak"
        seed = data / "seed.json"
        seed.write_text(json.dumps({"meta": {"month": "2026-10", "from": "seed"}, "tasks": [], "people": []}), encoding="utf-8")
        bak.write_text(json.dumps({"meta": {"month": "2026-10", "from": "bak", "rev": 3}, "tasks": [{"name": "DD-KEEP-1"}], "people": []}), encoding="utf-8")
        db_path.write_text("{not json", encoding="utf-8")
        old_db, old_seed = server.DB_PATH, server.SEED_PATH
        server.DB_PATH = db_path
        server.SEED_PATH = seed
        try:
            recovered = server._read_db_file()
            check("corrupt db.json falls back to bak", recovered.get("meta", {}).get("from") == "bak")
            check("bak keeps live tasks", (recovered.get("tasks") or [{}])[0].get("name") == "DD-KEEP-1")
            bak.write_text("also broken", encoding="utf-8")
            recovered = server._read_db_file()
            check("corrupt bak falls back to seed", recovered.get("meta", {}).get("from") == "seed")
            db_path.write_text("{still broken", encoding="utf-8")
            bak.unlink()
            seed.unlink()
            recovered = server._read_db_file()
            check("missing seed returns empty month db", recovered.get("meta", {}).get("month") == "2026-10")
        finally:
            server.DB_PATH = old_db
            server.SEED_PATH = old_seed


def test_save_db_atomic() -> None:
    with tempfile.TemporaryDirectory() as raw:
        data = Path(raw)
        db_path = data / "db.json"
        db_path.write_text(json.dumps({"meta": {"rev": 1}, "tasks": []}), encoding="utf-8")
        old = server.DB_PATH
        server.DB_PATH = db_path
        try:
            server.save_db({"meta": {"rev": 1, "ok": True}, "tasks": [{"name": "DD-X"}]})
            disk = json.loads(db_path.read_text(encoding="utf-8"))
            bak = json.loads(db_path.with_suffix(".json.bak").read_text(encoding="utf-8"))
            check("save_db writes the new db", disk.get("meta", {}).get("ok") is True)
            check("save_db keeps previous copy as bak", bak.get("meta", {}).get("rev") == 1)
            check("save_db bumps rev", int(disk.get("meta", {}).get("rev") or 0) >= 2)
        finally:
            server.DB_PATH = old


def test_employeers_pay_details() -> None:
    db = {
        "people": [
            {
                "id": "designer5",
                "name": "Олекса",
                "role": "designer",
                "payOrder": 6,
                "jobTitle": "Motion design",
                "legalName": "Злотко Олексій",
                "bankCard": "4323 3450 3785 1582",
                "iban": "ua94 3077 7000 0002 6208 7139 43056",
                "rateUsd": 15,
            }
        ]
    }
    check("stale Oleksa card is rewritten", finance_io.apply_bank_card_fixes({"people": [{"bankCard": "432345037851582"}]}))
    fixed = {"people": [{"bankCard": "432345037851582"}]}
    finance_io.apply_bank_card_fixes(fixed)
    check("stale Oleksa card becomes 16 digits", fixed["people"][0]["bankCard"] == "4323345037851582")
    check("spaces stripped from card", finance_io.normalize_bank_card("4323 3450 3785 1582") == "4323345037851582")
    check("iban compacted", finance_io.normalize_iban("ua94 3077700000026208713943056") == "UA943077700000026208713943056")
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "calc.xlsx"
        Workbook().save(path)
        finance_io.write_employeers(db, path)
        wb = load_workbook(path)
        ws = wb["Employeers"]
        cards = [ws.cell(r, 6).value for r in range(7, 16)]
        ibans = [ws.cell(r, 7).value for r in range(7, 16)]
        wb.close()
        check("Employeers writes Oleksa card as text", "4323345037851582" in [str(v) for v in cards if v], str(cards))
        check("Employeers writes Oleksa IBAN", "UA943077700000026208713943056" in [str(v) for v in ibans if v], str(ibans))


def test_playables_desk() -> None:
    class _C:
        pass

    cell = _C()
    fill = _C()
    fg = _C()
    fg.type = "theme"
    fg.theme = 8
    fg.tint = 0
    fill.patternType = "solid"
    fill.fgColor = fg
    cell.fill = fill
    theme_hex = excel_io._cell_fill_hex(cell)
    check("Google Sheets theme 8 is WIP blue", excel_io.status_from_line_fill(theme_hex) == "wip", theme_hex)
    check("DS-P name is Playables", direction_io.of_name("DS-P-060_MergeSolitaire") == direction_io.PLAYABLES)
    check("DD-AIV name is MGX", direction_io.of_name("DD-AIV-202-004_SettAI") == direction_io.MGX)
    purple = {"people": [{"id": "manager", "name": "Анастасія", "role": "manager", "color": "#9d50dd", "rateUsd": 5}]}
    direction_io.ensure_crew(purple)
    lp = next(p for p in purple["people"] if p["id"] == "manager")
    check("old LP purple becomes pastel orange", lp.get("color") == "#FFD9A8", str(lp.get("color")))
    db = {
        "meta": {"month": "2026-10", "clientRateUsd": 35, "playableClientRateUsd": 40, "taxPct": 0.06, "nbuRate": 41},
        "people": [
            {"id": "manager", "name": "Анастасія", "role": "manager", "rateUsd": 5, "directions": ["mgx", "playables"], "rates": {"mgx": 5, "playables": 7}},
            {"id": "artdir1", "name": "Анна", "role": "art_director", "rateUsd": 15, "directions": ["playables"], "rates": {"playables": 15}},
            {"id": "programmer1", "name": "Богдан", "role": "programmer", "rateUsd": 15, "directions": ["playables"], "rates": {"playables": 15}},
            {"id": "designer", "name": "Masha", "role": "designer", "rateUsd": 15, "directions": ["mgx"], "rates": {"mgx": 15}},
            {"id": "designer3", "name": "Сергій", "role": "designer", "rateUsd": 15},
        ],
        "tasks": [
            {
                "id": "mgx-1",
                "name": "DD-FIN-1",
                "month": "2026-10",
                "line": "DD",
                "department": "MGX",
                "hours": 9,
                "shifts": 1,
                "assigneeId": "designer",
                "help": [],
            },
            {
                "id": "pl-1",
                "name": "DS-P-060_MergeSolitaire",
                "month": "2026-10",
                "line": "DS",
                "department": "Playables",
                "hours": 30,
                "shifts": 0,
                "assigneeId": "artdir1",
                "assigneeHours": 10,
                "help": [],
                "programmerId": "programmer1",
                "programmerHours": 19,
                "producerHours": 29,
            },
        ],
    }
    direction_io.ensure_crew(db)
    sergey = next(p for p in db["people"] if p.get("id") == "designer3")
    check("Sergey sits Playables by default", "playables" in (sergey.get("directions") or []), str(sergey.get("directions")))
    check("Playables crew has Юля", any(p.get("id") == "artdir2" for p in db["people"]))
    check(
        "Playables Excel crew skips MOTION",
        all(p.get("role") != "designer" for p in playables_io.crew(db)),
        str([p.get("id") for p in playables_io.crew(db)]),
    )
    snap = finance_io.snapshot(db, "2026-10")
    labels = [r.get("label") for r in snap.get("byLine") or []]
    check("finance Client has MGX and Playable lines", any("MGX" in str(x) for x in labels) and any("Playable" in str(x) for x in labels), str(labels))
    play = next((r for r in snap["byLine"] if r.get("direction") == "playables"), None)
    check("Playable billed at $40", play and play.get("rate") == 40 and play.get("bill") == 1200, str(play))
    lp = next((r for r in snap["payroll"] if r.get("role") == "manager"), None)
    check(
        "LP Playables pay is producer hours × $7 not billed",
        lp and abs(float(lp.get("hoursPlay") or 0) - 29) < 0.05 and abs(float(lp.get("payUsd") or 0) - (9 * 5 + 29 * 7)) < 0.05,
        str(lp),
    )
    check("AD hours ignore billed buffer", excel_io.hours_for_person(db["tasks"][1], "artdir1") == 10)
    check("DEV hours sit on programmerId", excel_io.hours_for_person(db["tasks"][1], "programmer1") == 19)
    sticky = dict(db["tasks"][1])
    playables_io.attach_studio_buffer(sticky)
    check(
        "Excel leftover sticks as studioBuffer",
        abs(float(sticky.get("studioBuffer") or 0) - 1) < 0.05
        and abs(float(sticky.get("hours") or 0) - 30) < 0.05
        and abs(playables_io.crew_hours(sticky) - 29) < 0.05,
        str({k: sticky.get(k) for k in ("hours", "studioBuffer", "assigneeHours", "programmerHours")}),
    )
    sticky["programmerHours"] = 20
    playables_io.attach_studio_buffer(sticky)
    check(
        "studioBuffer stays when crew grows",
        abs(float(sticky.get("studioBuffer") or 0) - 1) < 0.05 and abs(float(sticky.get("hours") or 0) - 31) < 0.05,
        str({k: sticky.get(k) for k in ("hours", "studioBuffer", "programmerHours")}),
    )
    sticky["studioBuffer"] = 2
    playables_io.attach_studio_buffer(sticky)
    check(
        "DL studioBuffer is billed extra",
        abs(float(sticky.get("hours") or 0) - 32) < 0.05 and abs(playables_io.billed_hours(sticky) - 32) < 0.05,
        str(sticky.get("hours")),
    )
    stuck = {
        "id": "t-pl-devonly",
        "name": "DS-P-000_DevOnly",
        "department": "Playables",
        "hours": 5,
        "assigneeId": "programmer1",
        "assigneeHours": 0,
        "programmerId": "programmer1",
        "programmerHours": 4,
        "month": "2026-10",
    }
    board = {"people": list(db["people"]), "tasks": [dict(stuck)]}
    direction_io.ensure_crew(board)
    check(
        "normalize strips programmer from AD seat",
        playables_io.normalize_playable_seats(board) and not board["tasks"][0].get("assigneeId"),
        str(board["tasks"][0].get("assigneeId")),
    )
    check(
        "programmer-as-AD hours stay on DEV",
        playables_io.hours_for_person(board["tasks"][0], "programmer1") == 4
        and playables_io.hours_for_person(board["tasks"][0], "artdir1") == 0,
    )
    hours = excel_io.client_hours_bytes(db, "2026-10")
    wb = load_workbook(BytesIO(hours))
    names = [str(wb.active.cell(r, 1).value or "") for r in range(1, (wb.active.max_row or 1) + 1)]
    wb.close()
    check("hours xlsx has MGX and Playable sections", any(n.startswith("MGX ") for n in names) and any(n.startswith("Playable ") for n in names), str(names))
    check("hours xlsx splits MGX and Playable totals", any(n.startswith("MGX TOTAL") for n in names) and any(n.startswith("Playable TOTAL") for n in names), str(names[-8:]))
    rows = server.parse_email("DS-P-061-001_Stickerpuzzle 26\nDX-AAA-0001 0.3 shift")
    play_row = next(r for r in rows if r["name"].startswith("DS-P-"))
    mgx_row = next(r for r in rows if r["name"].startswith("DX-"))
    check("playable hours list is billed hours not shifts×9", play_row["hours"] == 26 and play_row["shifts"] == 0, str(play_row))
    check("MGX shift list still ×9", abs(mgx_row["hours"] - 2.7) < 0.05, str(mgx_row))
    if playables_io.XLSX_PATH.exists():
        parsed = playables_io.parse_workbook(playables_io.XLSX_PATH, "2026-10")
        names = [r["name"] for r in parsed.get("rows") or []]
        check("live Playables sheet parses DS-P rows", any(n.startswith("DS-P-") for n in names), str(names[:4]))
        scratch = {"meta": {"month": "2026-10"}, "people": list(db["people"]), "tasks": []}
        direction_io.ensure_crew(scratch)
        playables_io.apply_to_db(scratch, parsed, "2026-10")
        dice = next((t for t in scratch["tasks"] if "MemoryDice" in (t.get("name") or "")), None)
        merge = next((t for t in scratch["tasks"] if "MergeSolitaire" in (t.get("name") or "")), None)
        help_ids = {str((row or {}).get("id") or "") for row in ((dice or {}).get("help") or [])}
        check(
            "Playables ingest keeps lead AD id",
            bool(dice and dice.get("assigneeId") == "artdir1" and dice.get("programmerId") == "programmer1"),
            str({k: (dice or {}).get(k) for k in ("assigneeId", "assigneeHours", "programmerId", "programmerHours")}),
        )
        check(
            "Anna hours stay on 2DART 1",
            bool(dice and abs(float(dice.get("assigneeHours") or 0) - 1) < 0.05 and "artdir2" in help_ids),
            str({k: (dice or {}).get(k) for k in ("assigneeId", "assigneeHours", "help")}),
        )
        check(
            "Playables Excel WIP stays on the board as Ready",
            bool(dice and dice.get("studioStatus") == "new" and dice.get("excelStatus") == "wip"),
            str({k: (dice or {}).get(k) for k in ("studioStatus", "excelStatus")}),
        )
        wip_board = {"meta": {"month": "2026-10"}, "people": list(scratch["people"]), "tasks": []}
        direction_io.ensure_crew(wip_board)
        playables_io.apply_to_db(wip_board, parsed, "2026-10")
        held = next((t for t in wip_board["tasks"] if "MergeSolitaire" in (t.get("name") or "")), None)
        if held:
            held["studioStatus"] = "wip"
            held["mondayStatus"] = "WIP"
        playables_io.apply_to_db(wip_board, parsed, "2026-10")
        held = next((t for t in wip_board["tasks"] if "MergeSolitaire" in (t.get("name") or "")), None)
        check(
            "Playables ingest does not bounce board WIP back to Ready",
            bool(held and held.get("studioStatus") == "wip"),
            str({k: (held or {}).get(k) for k in ("studioStatus", "excelStatus", "mondayStatus")}),
        )
        check(
            "Anna keeps hours on art-only rows",
            bool(merge and merge.get("assigneeId") == "artdir1" and abs(float(merge.get("assigneeHours") or 0) - 9) < 0.05),
            str({k: (merge or {}).get(k) for k in ("assigneeId", "assigneeHours", "programmerId", "programmerHours")}),
        )
        kraken = next((t for t in scratch["tasks"] if "Kraken" in (t.get("name") or "")), None)
        check(
            "DEV-only playable keeps AD empty",
            bool(kraken and not kraken.get("assigneeId") and kraken.get("programmerId") == "programmer1"),
            str({k: (kraken or {}).get(k) for k in ("assigneeId", "assigneeHours", "programmerId", "programmerHours")}),
        )
        check(
            "DEV-only billed remainder is not AD hours",
            bool(kraken and abs(float(kraken.get("assigneeHours") or 0)) < 0.05 and abs(float(kraken.get("programmerHours") or 0) - 4) < 0.05),
            str({k: (kraken or {}).get(k) for k in ("hours", "assigneeHours", "programmerHours")}),
        )
        check(
            "DEV-only hours stay on the programmer",
            bool(kraken and playables_io.hours_for_person(kraken, "programmer1") == 4),
            str(playables_io.hours_for_person(kraken, "programmer1") if kraken else None),
        )
        ad_view = auth_io.view_state(scratch, {"role": "art_director", "email": "lisenbart.games@gmail.com"})
        ad_dice = next((t for t in (ad_view.get("tasks") or []) if "MemoryDice" in (t.get("name") or "")), None)
        ad_anna = next((p for p in (ad_view.get("people") or []) if p.get("id") == "artdir1"), None)
        check("AD view hides producerHours", ad_dice is not None and "producerHours" not in ad_dice)
        check("AD view hides pay rates", ad_anna is not None and "rates" not in ad_anna and "iban" not in ad_anna)
        dl_view = auth_io.view_state(scratch, {"role": "finance", "email": "dmytro@lisenbart.games"})
        dl_dice = next((t for t in (dl_view.get("tasks") or []) if "MemoryDice" in (t.get("name") or "")), None)
        check("DL view keeps producerHours", bool(dl_dice and dl_dice.get("producerHours")))
        check(
            "Excel leftover becomes DL studio buffer",
            bool(
                dice
                and abs(float(dice.get("studioBuffer") or 0) - 1) < 0.05
                and abs(float(dice.get("hours") or 0) - 26) < 0.05
                and abs(float(dice.get("producerHours") or 0) - 25) < 0.05
            ),
            str({k: (dice or {}).get(k) for k in ("hours", "studioBuffer", "producerHours", "assigneeHours", "programmerHours")}),
        )
        check(
            "Kraken Excel leftover is studio buffer not AD",
            bool(
                kraken
                and abs(float(kraken.get("studioBuffer") or 0) - 1) < 0.05
                and abs(float(kraken.get("hours") or 0) - 5) < 0.05
                and not kraken.get("assigneeId")
            ),
            str({k: (kraken or {}).get(k) for k in ("hours", "studioBuffer", "programmerHours", "assigneeHours")}),
        )
        check("DL view keeps studioBuffer", bool(dl_dice and abs(float(dl_dice.get("studioBuffer") or 0) - 1) < 0.05))
        lp_view = auth_io.view_state(scratch, {"role": "manager", "email": "anastasiia@lisenbart.games"})
        lp_dice = next((t for t in (lp_view.get("tasks") or []) if "MemoryDice" in (t.get("name") or "")), None)
        check("LP view keeps assigneeHours", bool(lp_dice and lp_dice.get("assigneeHours")))
        check("LP view hides producerHours", lp_dice is not None and "producerHours" not in lp_dice)
        check("LP view hides studioBuffer", lp_dice is not None and "studioBuffer" not in lp_dice)
        check(
            "LP playable hours are crew without buffer",
            bool(lp_dice and abs(float(lp_dice.get("hours") or 0) - 25) < 0.05),
            str((lp_dice or {}).get("hours")),
        )
        held_ad = float((held or {}).get("assigneeHours") or 0)
        held_buf = float((held or {}).get("studioBuffer") or 0)
        if held:
            held["assigneeHours"] = held_ad + 1
        playables_io.apply_to_db(wip_board, parsed, "2026-10")
        held = next((t for t in wip_board["tasks"] if "MergeSolitaire" in (t.get("name") or "")), None)
        check(
            "second Excel ingest does not rewrite portal seats or buffer",
            bool(
                held
                and held.get("studioStatus") == "wip"
                and abs(float(held.get("assigneeHours") or 0) - (held_ad + 1)) < 0.05
                and abs(float(held.get("studioBuffer") or 0) - held_buf) < 0.05
                and abs(float(held.get("hours") or 0) - playables_io.billed_hours(held)) < 0.05
            ),
            str({k: (held or {}).get(k) for k in ("assigneeHours", "hours", "studioBuffer", "studioStatus")}),
        )
        billed_zero = {
            "name": "DS-P-SPLIT-HELP",
            "department": "Playables",
            "hours": 0,
            "assigneeId": "artdir1",
            "assigneeHours": 8,
            "help": [{"id": "artdir2", "hours": 2}],
        }
        excel_io.normalize_help(billed_zero, {"artdir1", "artdir2"})
        check(
            "Playables help is not wiped when billed is 0",
            billed_zero.get("help") == [{"id": "artdir2", "hours": 2}]
            and abs(float(billed_zero.get("assigneeHours") or 0) - 8) < 0.05,
            str(billed_zero),
        )
        mgx_clip = {
            "name": "DD-MGX-CLIP",
            "department": "MGX",
            "hours": 0,
            "assigneeId": "designer",
            "help": [{"id": "designer2", "hours": 2}],
        }
        excel_io.normalize_help(mgx_clip, {"designer", "designer2"})
        check("MGX help still clips to billed hours", mgx_clip.get("help") == [], str(mgx_clip))
        js = (ROOT / "public" / "assets" / "app.js").read_text(encoding="utf-8")
        check("LP edits playable crew hours", "playableHoursHtml(t, canEditPlayableCrew())" in js)
        check("billed buffer is DL only", "const financeStrip = isFinance()" in js and "Studio buffer" in js)
        check("playable card has no Lead AD", "Lead AD" not in js and "Assign lead AD" not in js)
        check(
            "playable splits AD onto AD and DEV onto DEV",
            '"adSplitAdd"' in js and '"devSplitAdd"' in js and "programmerHelp" in js,
        )
        check("playable hours sit in AD and DEV blocks", "play-seat-${kind}" in js and "function playableSeatHtml" in js)
        check("DEV split picker is on the hours block", "— ${label}" in js and "pick DEV" in js)
        check(
            "Playables pick AD/DEV on hours seats",
            'class="play-owner-pick"' in js and "function takePlayableOwner" in js and "function deskHours" in js,
        )
        check(
            "pipeline Time uses crew hours, not billed",
            "playableCrewTotal(t)" in js.split("function timeCell")[1].split("function statusRank")[0]
            and "playableDeskTotal(t)" in js.split("function timeCell")[1].split("function statusRank")[0]
            and "Billed hours" not in js.split("function timeCell")[1].split("function statusRank")[0]
            and "hoursForPerson(t, me().id)" in js.split("function timeCell")[1].split("function statusRank")[0],
        )
        check(
            "Playables Excel Hours is crew, DL adds buffer",
            "isFinance() ? playableDeskTotal(t) : playableCrewTotal(t)"
            in js.split("function renderExcelSheet")[1].split("function statusMark")[0]
            and 'play ? `<th>Hours</th>`' in js,
        )
        check(
            "Playables Excel seats are LP / ARTD / DEV",
            "return `LP ${used.lp}`" in js and "return `ARTD ${used.art}`" in js and "return `DEV ${used.dev}`" in js,
        )
        check(
            "MGX Excel seats are LP / TL / MD",
            "return `TL ${used.tl}`" in js and "return `MD ${used.md}`" in js and "`EMPL ${i + 1}`" not in js,
        )
        check(
            "DL card edits studioBuffer and billed follows crew",
            'id="stStudioBuffer"' in js
            and "patch.studioBuffer" in js
            and 'id="playBilled"' in js
            and 'id="stHours"' not in js.split("function playableHoursHtml")[1].split("function readPlayableSplits")[0],
        )
        check("playable AD is never the programmer", "function playableAdId" in js and "if (id === String(t?.programmerId || \"\")) return \"\"" in js)
        check("playable hoursOwned is assigneeHours", "if (isPlayables(t)) return roundHours(t?.assigneeHours)" in js)
        check(
            "playable split writes remainder into owner hours",
            "applyPlayablePoolToOwner" in js
            and 'el.value = rest > 0 ? String(rest)' in js
            and "play-seat-keep" not in js
            and "keeps</span>" not in js
            and "Split hours cannot exceed the AD hours" in js,
        )
        check(
            "hours split has no duplicate remainder copy",
            "Extra AD/DEV" not in js
            and "Client stays" not in js
            and "play-seat-note" not in js
            and "not extra SuperPlay hours" not in js
            and " · Who" not in js,
        )
        check(
            "create skips Drive ensure while pack is queued",
            "driveQueued" in js and "if (state.driveQueued[t.id]) return;" in js,
        )
        check("hours split AD uses playableAdId", "const adId = playableAdId(t)" in js and "person(t.assigneeId)?.name || \"Unassigned\")" not in js.split("function playableHoursHtml")[1].split("function readPlayableSplits")[0])
        py = (ROOT / "server.py").read_text(encoding="utf-8")
        check(
            "playable crew save allows LP",
            'role not in {"finance", "manager", "teamlead"}' in py
            and 'data.pop("studioBuffer", None)' in py
            and "playables_io.attach_studio_buffer" in py,
        )
        check("save strips programmer from playable AD", "playables_io.ad_id" in py)
        update_sh = (ROOT / "deploy" / "update.sh").read_text(encoding="utf-8")
        check(
            "board-update rewrites MGX seat headers",
            "MGX check seats rewritten" in update_sh
            and "excel_io.write_month" in update_sh
            and "finance_io.write_month" in update_sh,
        )
        check(
            "board-update keeps live db and Playables Excel",
            "data/db.json" in update_sh
            and "data/SP_PLBL_check.xlsx" in update_sh
            and "SUPERPLAY_Playable_Calculator.xlsx" in update_sh,
        )
        drive_py = (ROOT / "drive_sync.py").read_text(encoding="utf-8")
        push_body = drive_py.split("def push_workbooks")[1].split("def _kind")[0]
        check(
            "Drive push writes Playables check like MGX",
            "SP_PLBL_check.xlsx" in push_body and "playableExcel" in push_body,
        )
        check(
            "Drive push writes Playables calculator like MGX",
            "SUPERPLAY_Playable_Calculator.xlsx" in push_body and "playableCalculator" in push_body,
        )
        check("Playables rename rewrites the Excel row", "playables_io.rewrite_existing_name" in py)
        check("LP/DL skip Art-ready DEV assign gate", "playable_may_assign_dev(trial, role)" in py)
        css = (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8")
        check("DS line is not Need Fixing burgundy", ".line-DS { background: #0d9488" in css and ".line-DS { background: #bb3354" not in css)
        check("client hours report splits MGX and Playable totals", "MGX TOTAL" in js and "Playable TOTAL" in js)
        check("hours table totals omit rates", "MGX TOTAL · $" not in js and "Playable TOTAL · $" not in js)
        check("hours report is two direction sheets", "sp-sheet-${kind}" in js and 'sheetHtml("mgx"' in js and 'sheetHtml("play"' in js)
        check("hours KPIs filter a direction", 'data-hours-dir="' in js and "state.hoursDir" in js)
        html = (ROOT / "public" / "index.html").read_text(encoding="utf-8")
        check(
            "DL SuperPlay hours stay in the cabinet",
            'id: "report", label: "SuperPlay"' in js
            and 'id="navClientSuperplay"' in html
            and "function isHoursView" in js,
        )
        check("client page is /superplay", "`${location.origin}/superplay`" in js or "origin}/superplay`" in js)
        check("client page title is SuperPlay", "SuperPlay — Lisenbart Board" in js and "Hours — Lisenbart Board" not in js)
        check("client tab is SuperPlay not Hours", '{ id: "report", label: "Hours" }' not in js)
        check("card warns if it changed while open", "This card changed while it was open" in js)
        check("card shows who is writing", "is on this card" in js and "function paintCardWatch" in js)
        check("presence ping is not a db save", '"/api/presence"' in js)
        check("hours tables share column widths", "table-layout: fixed" in css and "sp-col-num" in css and "sp-col-name" in js)
        check("hours totals use pastel ladder", "--sp-total: #b5c9dc" in css and "--sp-total-sum: #9bb5ce" in css)
        check("hours final numbers are darker", ".sp-dir-total td.sp-hours" in css and "background: var(--sp-total-sum)" in css)
        check("hours Excel classifies by name prefix", "def hours_direction" in (ROOT / "excel_io.py").read_text(encoding="utf-8"))
        check("pipeline hours cell is 5.5ch", ".hours-in {" in css and "5.5ch" in css)
        check("pipeline hours input takes a comma", 'deskNumInput("shift-in"' in js and 'type="text"' in js and "fmtDeskNum" in js)
        check("pipeline hours has step arrows", "num-step-up" in js and "num-step-down" in css)
        check("hours arrows save on pointerup not mousedown", "pointerup" in js and "pointerdown" in js and 'bindNumStep(el, commit, 1)' in js)
        check(
            "chrome drops slogan copy",
            "studio money" not in js
            and "Create task writes" not in js
            and "desk ${" not in js
            and "on one tab" not in js
            and "invoices are based on" not in js,
        )
        html = (ROOT / "public" / "index.html").read_text(encoding="utf-8")
        check(
            "identity sits under Sign out",
            'class="who-tools"' in html and 'class="who-id"' in html and 'id="roleHint"' in html,
        )
        check(
            "search has a clear control",
            "function searchHtml" in js
            and "search-clear" in js
            and "Clear search" in js
            and ".search-wrap" in css
            and '<input class="search" placeholder="Search project" />' not in js.split("function searchHtml")[0],
        )
        check(
            "client identity is SuperPlay logo",
            'id="boardLogo"' in html
            and 'src="assets/superplay.svg"' in html
            and (ROOT / "public" / "assets" / "superplay.svg").is_file()
            and "boardLogo.hidden = !clientHours" in js
            and "background: #000" not in css.split(".board-logo")[1].split("}")[0]
            and 'id="whoLogo"' not in html,
        )
        check(
            "client hours put the logo where the SuperPlay word was",
            "title.hidden = clientHours" in js
            and 'id="boardLogo"' in html.split('class="board-title"')[1].split('class="pill"')[0]
            and ".board-title[hidden]" in css
            and ".pill[hidden]" in css,
        )
        check(
            "SuperPlay logo stays on the title row",
            "html.hours-host .top-row" in css
            and "flex-wrap: nowrap" in css.split("html.hours-host .top-row")[1].split("}")[0]
            and "html.hours-host .top-end { width: auto; }" in css,
        )
        check("PIN Open is a form submit", 'id="loginForm"' in js and 'type="submit"' in js and "submitDoorLogin" in js)
        check("Create task name is one line", 'id="createTaskName"' in html and "<textarea id=\"createTaskName\"" not in html)
        check("Create task Enter submits", 'id="createTaskForm"' in html and 'type="submit"' in html.split('id="createTaskGo"')[0][-80:])
        check("Create task no longer needs Ctrl+Enter", "e.metaKey || e.ctrlKey" not in js.split("createTaskForm")[-1][:400])
        check("card Save is not bound to Enter", "saveTask" not in js.split('document.addEventListener("keydown"')[-1].split("probeStudioApi")[0] or "Enter" not in js.split('document.addEventListener("keydown"')[-1].split("probeStudioApi")[0])
        check("role is not a page subtitle", 'class="role-hint"' not in html)
        check("who label is the name only", "${p.name} · ${jobMeta(p).code}" not in js)
        check("AD/DEV hours blocks have a tint", ".play-seat-ad" in css and ".play-seat-dev" in css)
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "pl.xlsx"
            shutil.copy(playables_io.XLSX_PATH, path)
            playables_io.write_month(scratch, "2026-10", path)
            wb = load_workbook(path)
            ws = excel_io.find_sheet(wb, "2026-10")
            names = [str(ws.cell(1, c).value or "") for c in range(4, 10)]
            labels = [str(ws.cell(2, c).value or "") for c in range(4, 10)]
            wb.close()
            check(
                "write_month does not insert MOTION",
                "MOTION" not in labels and "Сергій" not in names,
                f"{names} {labels}",
            )
            hours_head = str(ws.cell(2, 3).value or "")
            joined = " ".join(labels)
            check(
                "Playables Excel columns are Hours / LP / ARTD / DEV",
                hours_head == "Hours"
                and "LP 1" in labels
                and "ARTD 1" in labels
                and "DEV 1" in labels
                and "Producer" not in labels
                and "2DART" not in joined
                and "PRGMM" not in joined,
                f"{hours_head} {labels}",
            )
            parsed2 = playables_io.parse_workbook(path, "2026-10")
            dice2 = next((r for r in (parsed2.get("rows") or []) if "MemoryDice" in (r.get("name") or "")), None)
            check(
                "Anna hours survive write_month",
                bool(dice2 and abs(float(dice2.get("assigneeHours") or 0) - 1) < 0.05),
                str(dice2),
            )
            portal_task = {
                "id": "t-pl-portal",
                "name": "DS-P-999_PortalOnly",
                "line": "DS",
                "department": "Playables",
                "month": "2026-10",
                "hours": 3,
                "studioBuffer": 1,
                "assigneeId": "artdir1",
                "assigneeHours": 3,
                "programmerId": "programmer1",
                "programmerHours": 0,
                "studioStatus": "new",
            }
            playables_io.attach_studio_buffer(portal_task)
            scratch.setdefault("tasks", []).append(portal_task)
            playables_io.write_month(scratch, "2026-10", path)
            parsed3 = playables_io.parse_workbook(path, "2026-10")
            portal_row = next((r for r in (parsed3.get("rows") or []) if "PortalOnly" in (r.get("name") or "")), None)
            check(
                "portal Playables land on the Excel sheet",
                bool(
                    portal_row
                    and abs(float(portal_row.get("hours") or 0) - 4) < 0.05
                    and abs(float(portal_row.get("assigneeHours") or 0) - 3) < 0.05
                ),
                str(portal_row),
            )
            scratch["droppedNames"] = ["DS-P-999_PortalOnly"]
            scratch["tasks"] = [t for t in scratch["tasks"] if t.get("id") != "t-pl-portal"]
            playables_io.write_month(scratch, "2026-10", path)
            parsed4 = playables_io.parse_workbook(path, "2026-10")
            gone = next((r for r in (parsed4.get("rows") or []) if "PortalOnly" in (r.get("name") or "")), None)
            check("dropped Playables leave the Excel sheet", gone is None, str(gone))
        check(
            "Producer hours land on Anastasiia",
            playables_io.hours_for_person(dice, "manager") > 0,
            str(playables_io.hours_for_person(dice, "manager") if dice else None),
        )
        lp = next((r for r in snap["payroll"] if r.get("role") == "manager"), None)
        check(
            "Employeers pay splits MGX vs Playables",
            lp and float(lp.get("payUsdMgx") or 0) > 0 and float(lp.get("payUsdPlay") or 0) > 0,
            str({k: (lp or {}).get(k) for k in ("payUsdMgx", "payUsdPlay", "hoursMgx", "hoursPlay")}),
        )
        with tempfile.TemporaryDirectory() as raw:
            calc_path = Path(raw) / "play-calc.xlsx"
            play = dict(db["tasks"][1])
            playables_io.attach_studio_buffer(play)
            calc_db = {"meta": dict(db.get("meta") or {}), "people": list(db["people"]), "tasks": [play]}
            finance_io.write_playable_month(calc_db, "2026-10", path=calc_path)
            wb = load_workbook(calc_path)
            ws = wb[finance_io.sheet_name_for("2026-10")]
            heads = [str(ws.cell(2, c).value or "") for c in range(1, 16)]
            row_n = next(
                r
                for r in range(5, (ws.max_row or 5) + 1)
                if "MergeSolitaire" in str(ws.cell(r, 2).value or "")
            )
            hours = ws.cell(row_n, 4).value
            lp_col = next((c for c, n in enumerate(heads, 1) if n == "Анастасія"), 6)
            anna_col = next((c for c, n in enumerate(heads, 1) if n == "Анна"), 0)
            client_col = next((c for c, n in enumerate(heads, 1) if str(ws.cell(5, c).value or "") == "CLIENT IN"), 0)
            lp_formula = str(ws.cell(row_n, lp_col).value or "")
            anna_formula = str(ws.cell(row_n, anna_col).value or "") if anna_col else ""
            client_formula = str(ws.cell(row_n, client_col).value or "") if client_col else ""
            wb.close()
            seats = [str(ws.cell(5, c).value or "") for c in range(6, 12)]
            check(
                "Playables calculator seats are LP / ARTD / DEV",
                "LP 1" in seats and "ARTD 1" in seats and "EMPL 1" not in seats,
                str(seats),
            )
            check(
                "Playables calculator HOURS are crew without buffer",
                abs(float(hours or 0) - 29) < 0.05,
                str(hours),
            )
            check(
                "Playables calculator LP pay is HOURS cell × rate",
                lp_formula.replace(" ", "") == f"=D{row_n}*7",
                lp_formula,
            )
            check(
                "Playables calculator CLIENT IN is HOURS cell × $40",
                client_formula.replace(" ", "") == f"=D{row_n}*40",
                client_formula,
            )
            check(
                "Playables calculator AD is hours × rate",
                "*15" in anna_formula.replace(" ", "") and "10" in anna_formula,
                anna_formula,
            )


def test_playable_card_pipeline() -> None:
    leftover = {"studioStatus": "storyboard", "assigneeId": "artdir1", "department": "Playables"}
    check("old Storyboard chip is Check", direction_io.playable_norm_status("storyboard") == "check")
    check("old HTML ready chip is Done", direction_io.playable_norm_status("htmlready") == "done")
    check("old Build chip is Art ready", direction_io.playable_norm_status("build") == "artready")
    check("Quote leftover is Ready", direction_io.playable_norm_status("quote") == "new")
    check("WIP stays WIP on Playables", direction_io.playable_norm_status("wip") == "wip")
    quote = {"studioStatus": "quote", "assigneeId": "artdir1", "programmerId": "programmer1"}
    check(
        "Quote leftover AD starts with WIP",
        direction_io.playable_next_statuses("art_director", quote) == ["wip"],
    )

    ad = {"studioStatus": "new", "assigneeId": "artdir1"}
    check("AD Ready only offers WIP", direction_io.playable_next_statuses("art_director", ad) == ["wip"])
    check("DEV work still waits before Art ready", not direction_io.playable_may_assign_dev(ad))
    check("LP may assign DEV on Ready", direction_io.playable_may_assign_dev(ad, "manager"))
    check("DL may assign DEV on Ready", direction_io.playable_may_assign_dev(ad, "finance"))
    check("DEV with AD waits on Ready", direction_io.playable_next_statuses("programmer", {**ad, "programmerId": "programmer1"}) == [])
    ad["studioStatus"] = "wip"
    check("AD WIP offers Check", direction_io.playable_next_statuses("art_director", ad) == ["check"])
    art_wip = {"studioStatus": "wip", "assigneeId": "artdir1", "programmerId": "programmer1"}
    check("DEV on the card is not HTML yet", not direction_io.playable_in_html(art_wip))
    check(
        "AD WIP with DEV still offers Check",
        direction_io.playable_next_statuses("art_director", art_wip) == ["check"],
    )
    check(
        "DEV on AD WIP has no HTML chips",
        direction_io.playable_next_statuses("programmer", art_wip) == [],
    )
    ad["studioStatus"] = "check"
    check("AD waits on Check for LP", direction_io.playable_next_statuses("art_director", ad) == [])
    check("LP does not chip-set Send", not direction_io.playable_can_set("manager", ad, "approve")[0])
    check("LP can Mark Closed from Check", direction_io.playable_can_set("manager", ad, "closed")[0])
    lp_from_ready = {"studioStatus": "new", "assigneeId": "artdir1", "department": "Playables"}
    for key in ("wip", "check", "revision", "finalart", "artready", "done", "closed"):
        ok, err = direction_io.playable_can_set("manager", lp_from_ready, key)
        check(f"LP can set Playables {key}", ok, err)
        ok_dl, err_dl = direction_io.playable_can_set("finance", lp_from_ready, key)
        check(f"DL can set Playables {key}", ok_dl, err_dl)
    check(
        "LP Playables chips include the work pipeline",
        set(direction_io.playable_next_statuses("manager", lp_from_ready))
        >= {"new", "wip", "check", "revision", "finalart", "artready", "done", "closed"},
    )

    bounced = {"studioStatus": "clientart", "assigneeId": "artdir1"}
    check(
        "LP Sent bounce or approve board",
        set(direction_io.playable_next_statuses("manager", bounced)) >= {"revision", "finalart"},
    )
    bounced["studioStatus"] = "revision"
    check("AD Need Fixing offers WIP", direction_io.playable_next_statuses("art_director", bounced) == ["wip"])
    check("DEV stays off storyboard Fix", direction_io.playable_next_statuses("programmer", bounced) == [])

    ad["studioStatus"] = "finalart"
    check("AD Final art offers Art ready", direction_io.playable_next_statuses("art_director", ad) == ["artready"])
    helper_ok, helper_err = direction_io.playable_can_set("art_director", ad, "artready", actor_id="artdir2")
    check("helper AD cannot set status", not helper_ok, helper_err)
    ad["studioStatus"] = "artready"
    check("DEV assign unlocks on Art ready", direction_io.playable_may_assign_dev(ad))
    check("DEV on Art ready offers WIP", direction_io.playable_next_statuses("programmer", {**ad, "programmerId": "programmer1"}) == ["wip"])
    check("AD is done at Art ready", direction_io.playable_next_statuses("art_director", ad) == [])

    html = {
        "studioStatus": "wip",
        "assigneeId": "artdir1",
        "programmerId": "programmer1",
        "playableHtml": True,
    }
    check("HTML loop starts after Art ready", direction_io.playable_in_html(html))
    check("DEV HTML WIP offers Done", direction_io.playable_next_statuses("programmer", html) == ["done"])
    check("AD does not drive HTML WIP", direction_io.playable_next_statuses("art_director", html) == [])
    ok_html, err_html = direction_io.playable_can_set("programmer", html, "done", actor_id="programmer1")
    check("DEV can pick Done from HTML WIP", ok_html, err_html)
    html_bounce = {
        "studioStatus": "approve",
        "assigneeId": "artdir1",
        "programmerId": "programmer1",
        "playableHtml": True,
    }
    check("LP HTML bounce is Need Fixing", "revision" in direction_io.playable_next_statuses("manager", html_bounce))
    html_bounce["studioStatus"] = "revision"
    check("DEV HTML Fix offers WIP", direction_io.playable_next_statuses("programmer", html_bounce) == ["wip"])
    check("AD does not take HTML Fix", direction_io.playable_next_statuses("art_director", html_bounce) == [])

    dev_only = {"studioStatus": "new", "programmerId": "programmer1"}
    check("DEV-only Ready offers WIP", direction_io.playable_next_statuses("programmer", dev_only) == ["wip"])
    check("DEV-only can be assigned from Ready", direction_io.playable_may_assign_dev(dev_only))

    check("MD has no Playables status", direction_io.playable_next_statuses("designer", ad) == [])
    check("TL has no Playables status", direction_io.playable_next_statuses("teamlead", ad) == [])
    split = {
        "assigneeId": "artdir1",
        "programmerId": "programmer1",
        "programmerHelp": [{"id": "programmer2", "hours": 1}],
    }
    check("DEV split is on the playable seat", direction_io.on_playable_seat(split, "programmer2"))
    js = (ROOT / "public" / "assets" / "app.js").read_text(encoding="utf-8")
    check("card uses next-step Playables statuses", "function playableNextStatuses" in js and "function playableStatusActions" in js)
    check("Playables assign is LP/DL", "function canAssignPlayable" in js and "play ? canAssignPlayable() : canAssign()" in js)
    check("LP seats DEV from Ready in copy", "You can already put DEV on the card" in js)
    check("Playables guide names the pipeline", "Ready → WIP → Check → Sent → Final art → Art ready → WIP → Done → Send → Closed" in js)
    check("AD owns Playables status copy", "AD owns status — you do not change it" in js)
    check(
        "LP/DL can set any Playables work status",
        'return PLAY_OPS_STATUSES.slice()' in js.split("function playableNextStatuses")[1].split("function playableStatusActions")[0]
        and "If they are away, set the status yourself" in js,
    )
    check("DEV HTML paste is owner-only", "isProgrammer() && String(t.programmerId || \"\") === String(me().id) && playableInHtml(t)" in js)
    check("AD can paste a storyboard link", "isArtDirector() && playableAdId(t) === me().id && !playableInHtml(t)" in js)
    check("empty Status plaque is hidden", "const statusKeys = statusActions(t)" in js and "statusKeys.length" in js)
    check("Check needs the storyboard link", '(key === "done" || key === "check")' in js)
    check("Check unlocks after the storyboard paste", '["tlcheck", "check", "done"].forEach' in js)
    check("playableInHtml uses Art ready flag", "t?.playableHtml" in js)
    check("designer Playables status set is empty", "designer: []" in js.split("const PLAY_STATUS_BY_ROLE")[1].split("const LINES")[0])

    ready_check = {
        "department": "Playables",
        "studioStatus": "check",
        "assigneeId": "artdir1",
        "resultUrl": "",
    }
    check("Check without storyboard is not sendable", not server.is_sendable(ready_check))
    ready_check["resultUrl"] = "https://board.example/story"
    check("Check with storyboard is sendable", server.is_sendable(ready_check))
    art = {"studioStatus": "artready", "assigneeId": "artdir1", "programmerId": "programmer1", "resultUrl": "https://board.example/story"}
    direction_io.touch_playable_html(art)
    check("Art ready locks the HTML loop", art.get("playableHtml") is True)
    check("Art ready keeps the storyboard URL", art.get("storyboardUrl") == "https://board.example/story")


def test_card_presence() -> None:
    server._PRESENCE.clear()
    server.touch_presence("tok-a", {"personId": "manager", "role": "manager"}, "t-card")
    other = server.presence_view("tok-b")
    mine = server.presence_view("tok-a")
    check(
        "presence lists the other session",
        any(not r.get("self") for r in other.get("t-card") or []),
        str(other),
    )
    check(
        "presence marks self",
        any(r.get("self") for r in mine.get("t-card") or []),
        str(mine),
    )
    server.drop_presence("tok-a")
    check("presence clears on leave", "t-card" not in server.presence_view("tok-a"))
    src = (ROOT / "server.py").read_text(encoding="utf-8")
    check("presence does not bump db rev", "def touch_presence" in src and "bump_rev" not in src.split("def touch_presence")[1].split("def drop_presence")[0])


def test_urgent_marker() -> None:
    src = (ROOT / "server.py").read_text(encoding="utf-8")
    js = (ROOT / "public" / "assets" / "app.js").read_text(encoding="utf-8")
    css = (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8")
    check("pipeline header is Dead Line", '<div class="cell">Dead Line</div>' in js)
    check("Due is no longer the pipeline header", '<div class="cell">Due</div>' not in js)
    check("urgent is not a studio status", '"urgent"' not in src.split("STATUS_BY_ROLE")[0] and "urgent" not in server.STATUS_BY_ROLE.get("manager", []))
    check("LP/DL may mark urgent", "only manager or finance can mark urgent" in src)
    check("card has URGENT toggle", 'id="stUrgent"' in js and "function urgentMark" in js)
    check("urgent mark is a ! badge", 'class="urgent-mark"' in js and ".urgent-mark" in css)
    check("urgent sorts above same status", "Number(isUrgent(b)) - Number(isUrgent(a))" in js)
    closed_bit = src.split('elif nxt == "new":')[-1].split("note = data.get")[0]
    check("Closed drops URGENT on the server", 'task["urgent"] = False' in closed_bit)
    check("Mark Closed sends urgent off", "urgent: false" in js.split('$("closeOne")')[-1].split("fillDrivePackIfMissing")[0])


def test_card_folder_buttons() -> None:
    js = (ROOT / "public" / "assets" / "app.js").read_text(encoding="utf-8")
    css = (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8")
    check("card Open folder is now Task folder", ">Task folder</a>" in js)
    check("card no longer says Open folder", "Open folder" not in js)
    check("card always paints Working folder", ">Working folder</a>" in js and ">Working folder</button>" in js)
    check("idle Working folder is a disabled button", 'class="btn pack-work" disabled' in js)
    check("Working folder reads /drive/u/N/links", r"u\/\d+" in js)
    check("studio Drive pack is Task folder", "function taskFolder" in js and "function hasDriveFolder" in js)
    check("Working folder has its own chip style", ".pack-work" in css)
    sample = (
        "Working folder:\n"
        "https://drive.google.com/drive/u/1/folders/1y5WRZSCqcN9rHRHSkO4WS-oOmil-Pozr"
    )
    found = re.search(
        r"(?:working folder|pls work here|please work here|work here|pls work)[:\s]*\n*"
        r"(https://drive\.google\.com/(?:drive/(?:u/\d+/)?folders/|open\?id=)[a-zA-Z0-9_-]+)",
        sample,
        re.I,
    )
    check(
        "Working folder line with /drive/u/1/ is a Drive folder",
        bool(found and "/folders/1y5WRZSCqcN9rHRHSkO4WS-oOmil-Pozr" in found.group(1)),
        str(found.group(1) if found else None),
    )


def test_superplay_visits_and_status() -> None:
    db = {
        "meta": {"month": "2026-10", "clientRateUsd": 35},
        "people": [
            {"id": "client", "name": "SuperPlay", "role": "client"},
            {"id": "finance", "name": "Dmytro", "role": "finance", "iban": "UA1"},
        ],
        "tasks": [
            {
                "id": "t-1",
                "name": "DD-A",
                "line": "DD",
                "month": "2026-10",
                "hours": 9,
                "studioStatus": "wip",
                "assigneeId": "md",
                "brief": "secret",
            }
        ],
        "events": [{"text": "nope"}],
    }
    client = auth_io.client_state(db)
    task = (client.get("tasks") or [{}])[0]
    check("client hours carry studio status", task.get("studioStatus") == "wip")
    check("client hours hide assignee and brief", "assigneeId" not in task and "brief" not in task)
    check("client state has no visit log", "superplayVisits" not in client)
    fin = auth_io.view_state(db, {"role": "finance", "email": "dmytro@lisenbart.games"})
    check("finance state has visit log", isinstance(fin.get("superplayVisits"), list))
    check("finance visit log is not written into db", "superplayVisits" not in db)
    prev = auth_io.VISITS_PATH
    tmp = Path(tempfile.mkdtemp()) / "visits.json"
    auth_io.VISITS_PATH = tmp
    try:
        auth_io.record_client_visit(login="Olena", ip="10.0.0.8", ua="Mozilla/5.0 Chrome/120.0")
        rows = auth_io.client_visits()
        check(
            "sign-in log keeps login and ip",
            bool(rows) and rows[0].get("login") == "Olena" and rows[0].get("ip") == "10.0.0.8",
            str(rows[:1]),
        )
        check("sign-in log has a timestamp", bool((rows or [{}])[0].get("at")))
        check(
            "private IP is marked local, not a city",
            (rows or [{}])[0].get("isp") == "local" and not (rows or [{}])[0].get("country"),
            str(rows[:1]),
        )
        play = auth_io.place_from_host("213-134-163-199.dynamic.play.pl")
        check(
            "Play.pl reverse DNS is Poland / Play",
            play.get("country") == "Poland" and play.get("isp") == "Play",
            str(play),
        )
        kyivstar = auth_io.place_from_host("host.kyivstar.net")
        check(
            "Kyivstar reverse DNS is Ukraine",
            kyivstar.get("country") == "Ukraine" and kyivstar.get("isp") == "Kyivstar",
            str(kyivstar),
        )
        bezeq = auth_io.place_from_host("bzq-1.static.bezeqint.net")
        check(
            "Bezeq reverse DNS is Israel",
            bezeq.get("country") == "Israel" and bezeq.get("isp") == "Bezeq",
            str(bezeq),
        )
        old = [{"at": "2026-10-18T12:29:00Z", "login": "superplay", "ip": "10.0.0.8", "ua": "Chrome"}]
        tmp.write_text(json.dumps(old), encoding="utf-8")
        filled = auth_io.client_visits()
        check(
            "old sign-in rows get a place on read",
            bool(filled) and filled[0].get("isp") == "local",
            str(filled[:1]),
        )
    finally:
        auth_io.VISITS_PATH = prev
    js = (ROOT / "public" / "assets" / "app.js").read_text(encoding="utf-8")
    py = (ROOT / "server.py").read_text(encoding="utf-8")
    check("SuperPlay table has STATUS", ">STATUS</th>" in js and "${statusChip(t)}" in js)
    check(
        "DL SuperPlay tab shows sign-ins",
        "function superplayVisitsHtml" in js and "isFinance() ? superplayVisitsHtml()" in js,
    )
    check(
        "sign-in Where is IP country, not Kyiv on the clock",
        "function visitPlace" in js
        and ">Where</th>" in js
        and "Clock is Kyiv time" in js
        and "(Kyiv)" not in js.split("function kyivVisitStamp")[1].split("function uaLabel")[0],
    )
    check(
        "SuperPlay Login is a name, PIN is the key",
        'if email not in {auth_io.CLIENT_EMAIL, "superplay@lg-board"}:' not in py,
    )


def test_lp_can_close_without_send() -> None:
    src = (ROOT / "server.py").read_text(encoding="utf-8")
    js = (ROOT / "public" / "assets" / "app.js").read_text(encoding="utf-8")
    check("LP role may set closed", "closed" in server.STATUS_BY_ROLE["manager"] and "closed" in server.STATUS_BY_ROLE["finance"])
    check(
        "LP/DL may set any MGX work status",
        {"new", "wip", "revision", "tlcheck", "done", "closed"} <= set(server.STATUS_BY_ROLE["manager"])
        and {"new", "wip", "revision", "tlcheck", "done", "closed"} <= set(server.STATUS_BY_ROLE["finance"]),
    )
    check(
        "MGX chips for LP/DL are the work pipeline",
        "const MGX_OPS_STATUSES" in js and "return MGX_OPS_STATUSES.slice()" in js.split("function statusActions")[1].split("function statusButtonDisabled")[0],
    )
    check(
        "MD on TL Check does not get WIP",
        'if (cur === "tlcheck") return ["tlcheck"]' in js.split("function statusActions")[1].split("function statusButtonDisabled")[0],
    )
    check(
        "status chip writes without waiting for Save",
        "commitCard({ fromChip: true, status: key })" in js and "close: !fromChip" in js,
    )
    check(
        "dirty Save is a green ready chip",
        'id="saveTask"' in js
        and "card-save" in js
        and "is-ready" in js
        and "#saveTask.card-save.is-ready" in (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8")
        and "justify-content: center" in (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8").split("#saveTask.card-save")[1].split("}")[0]
        and "appearance: none" in (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8").split("#saveTask.card-save")[1].split("}")[0]
        and "line-height: 1" in (ROOT / "public" / "assets" / "board.css").read_text(encoding="utf-8").split("#saveTask.card-save")[1].split("}")[0],
    )
    check(
        "close asks Save not discard",
        'id="cardSaveModal"' in (ROOT / "public" / "index.html").read_text(encoding="utf-8")
        and "Don’t save" in (ROOT / "public" / "index.html").read_text(encoding="utf-8")
        and "Discard unsaved changes?" not in js,
    )
    check(
        "Who works this writes at once",
        'commitCard({ fromChip: true })' in js
        and "writeAssign" in js,
    )
    check(
        "Save compresses screenshots and waits for the note",
        "function canvasShot" in js
        and "MAX_NOTE_BYTES" in js
        and "timeout: 90000" in js
        and "noteError" in js
        and "Saving ${n} screenshot" in js
        and "applyState(out.state, { render: reopen !== false })" in js,
    )
    check(
        "card shows who set the status and when",
        "function statusLogHtml" in js
        and "status-log" in js
        and "whenOf" in (ROOT / "public" / "assets" / "notices.js").read_text(encoding="utf-8")
        and "task.wip" in (ROOT / "public" / "assets" / "notices.js").read_text(encoding="utf-8"),
    )
    check("LP/DL skip Done-after-TL-Check", "and role not in OPS_ROLES" in src.split("Done is after TL Check")[0][-250:])
    check("server no longer requires Send before Closed", "Closed is after Send" not in src)
    check("card offers Closed without approve-only gate", "Closed is after Send" not in js)
    check("Mark Closed is not locked to Appr.", 't.studioStatus === "approve" ? `<button type="button" class="btn" id="closeOne">' not in js)


def test_finance_reads_board_not_xlsx() -> None:
    called = {"n": 0}
    orig = excel_io.parse_workbook

    def boom(*_a, **_k):
        called["n"] += 1
        raise AssertionError("finance must not parse Excel")

    excel_io.parse_workbook = boom
    try:
        db = {
            "meta": {"month": "2026-10", "nbuRate": 41, "clientRateUsd": 35, "taxPct": 0.06},
            "people": [{"id": "designer", "name": "Masha", "role": "designer", "rateUsd": 15}],
            "tasks": [
                {
                    "id": "t-1",
                    "name": "DD-FIN-1",
                    "month": "2026-10",
                    "line": "DD",
                    "hours": 9,
                    "shifts": 1,
                    "assigneeId": "designer",
                    "help": [],
                }
            ],
        }
        rows = excel_io.tasks_for_finance(db, "2026-10")
        check("finance hours come from the board", len(rows) == 1 and called["n"] == 0, str(rows))
        snap = finance_io.snapshot(db, "2026-10")
        check("snapshot does not parse Excel", called["n"] == 0, str(called))
        check("snapshot hours match the board", snap.get("hours") == 9, str(snap.get("hours")))
    finally:
        excel_io.parse_workbook = orig


def test_status_receipt() -> None:
    import notice_io

    db = {
        "meta": {},
        "people": [
            {"id": "designer3", "name": "Сергій", "role": "designer"},
            {"id": "teamlead", "name": "Настя", "role": "teamlead"},
            {"id": "manager", "name": "Анастасія", "role": "manager"},
        ],
        "notices": [],
        "noticeReads": {},
        "tasks": [],
    }
    task = {"id": "t-1", "name": "DX-VEO-065-003", "assigneeId": "designer3", "studioStatus": "tlcheck"}
    notice_io.on_task_change(db, task, {"studioStatus": "wip"}, "designer3")
    row = (task.get("statusLog") or [None])[-1] or {}
    check(
        "TL Check stamps who and when",
        row.get("to") == "tlcheck" and row.get("actorName") == "Сергій" and row.get("at"),
        str(row),
    )
    check(
        "TL Check pings TL",
        any(n.get("kind") == "task.tlcheck" and "teamlead" in (n.get("audience") or []) for n in db.get("notices") or []),
        str(db.get("notices")),
    )
    task["studioStatus"] = "wip"
    notice_io.on_task_change(db, task, {"studioStatus": "tlcheck"}, "designer3")
    back = (task.get("statusLog") or [None])[-1] or {}
    check("WIP after TL Check is on the card log", back.get("to") == "wip" and len(task.get("statusLog") or []) == 2, str(task.get("statusLog")))
    ping = next((n for n in db.get("notices") or [] if n.get("kind") == "task.wip"), None)
    check(
        "WIP after TL Check pings TL and LP, not the actor",
        ping is not None
        and "teamlead" in (ping.get("audience") or [])
        and "manager" in (ping.get("audience") or [])
        and "designer3" not in (ping.get("audience") or []),
        str(ping),
    )
    task["studioStatus"] = "wip"
    notice_io.on_task_change(db, task, {"studioStatus": "new"}, "designer3")
    check(
        "normal WIP from Ready does not ping Back to WIP",
        sum(1 for n in db.get("notices") or [] if n.get("kind") == "task.wip") == 1,
        str([n.get("kind") for n in db.get("notices") or []]),
    )


def test_comment_shots() -> None:
    orig = server.COMMENT_FILES
    tmp = Path(tempfile.mkdtemp())
    server.COMMENT_FILES = tmp
    try:
        db = {"meta": {"commentSeq": 0}, "people": [{"id": "manager", "name": "A"}]}
        task = {"id": "t-test-shot", "name": "DD-X-1", "comments": []}
        bad = server.make_comment(
            db,
            actor_id="manager",
            body="hi",
            extra=[{"mime": "image/png", "data": "!!!!"}],
            task=task,
        )
        check("unreadable screenshots fail the note", bad == "could not read those screenshots", str(bad))
        tiny = (
            "data:image/png;base64,"
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        )
        ok = server.make_comment(
            db,
            actor_id="manager",
            body="logo left",
            extra=[{"mime": "image/png", "data": tiny}],
            task=task,
        )
        check(
            "tiny screenshot stores a file",
            isinstance(ok, dict) and any((a.get("kind") == "image") for a in ok.get("attachments") or []),
            str(ok),
        )
    finally:
        server.COMMENT_FILES = orig
        shutil.rmtree(tmp, ignore_errors=True)
    check(
        "payload cap names screenshots",
        "screenshots are too large together" in (ROOT / "server.py").read_text(encoding="utf-8"),
    )


def test_sync_files_releases_lock() -> None:
    original = server.sync_excel
    during = {"locked": None}

    def slow(db, month=None, calculator=False):
        time.sleep(0.4)
        during["locked"] = server.LOCK.locked()
        return {"ok": False, "skipped": True}

    server.sync_excel = slow
    try:
        th = threading.Thread(target=lambda: server._sync_files_bg("2026-10", False), daemon=True)
        th.start()
        time.sleep(0.08)
        t0 = time.monotonic()
        got = server.LOCK.acquire(timeout=0.25)
        dt = time.monotonic() - t0
        if got:
            server.LOCK.release()
        th.join(timeout=2)
        check(
            "excel bg does not hold board LOCK",
            got is True and dt < 0.2,
            f"got={got} dt={dt:.3f}s",
        )
        check("LOCK.locked() is false during xlsx/Drive", during["locked"] is False, str(during))
    finally:
        server.sync_excel = original


def test_unreadable_ingest_does_not_wipe() -> None:
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / "bad.xlsx"
        path.write_bytes(b"")
        parsed = excel_io.parse_workbook(path, "2026-10")
        db = _db()
        db["tasks"] = [{"id": "t-1", "name": "DD-KEEP-9", "month": "2026-10", "shifts": 1, "hours": 9}]
        before = list(db["tasks"])
        if not parsed.get("sheet"):
            check("empty parse does not apply", True)
            check("existing cards stay when Excel is unreadable", db["tasks"] == before)
        else:
            check("empty parse does not apply", False, "sheet was set")


def _http(url: str, data=None, cookie: str | None = None, timeout: float = 8):
    headers = {"Content-Type": "application/json"}
    if cookie:
        headers["Cookie"] = cookie
    body = None if data is None else (data if isinstance(data, bytes) else json.dumps(data).encode("utf-8"))
    req = urllib.request.Request(url, data=body, headers=headers, method="POST" if data is not None else "GET")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as res:
            raw = res.read()
            ctype = res.headers.get("Content-Type") or ""
            payload = json.loads(raw.decode("utf-8")) if "json" in ctype else raw
            return res.status, payload, res.headers
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            payload = json.loads(raw.decode("utf-8"))
        except Exception:
            payload = raw.decode("utf-8", errors="replace")
        return e.code, payload, e.headers


def test_http(base: str) -> str | None:
    code, payload, _ = _http(f"{base}/health")
    check("GET /health", code == 200 and payload.get("ok") is True and payload.get("build"), str(payload))
    code, payload, _ = _http(f"{base}/api/version")
    check("GET /api/version", code == 200 and payload.get("build") == server.BUILD, str(payload.get("build")))

    code, payload, headers = _http(f"{base}/favicon.ico")
    ico_type = (headers.get("Content-Type") or "").lower()
    ico_cache = (headers.get("Cache-Control") or "").lower()
    ico_ok = (
        code == 200
        and isinstance(payload, (bytes, bytearray))
        and bytes(payload[:4]) == b"\x00\x00\x01\x00"
        and "icon" in ico_type
        and "no-store" not in ico_cache
    )
    check("GET /favicon.ico is a cacheable icon", ico_ok, f"{code} {ico_type} {ico_cache}")
    code, payload, headers = _http(f"{base}/apple-touch-icon.png")
    png_type = (headers.get("Content-Type") or "").lower()
    png_cache = (headers.get("Cache-Control") or "").lower()
    png_ok = (
        code == 200
        and isinstance(payload, (bytes, bytearray))
        and bytes(payload[:8]) == b"\x89PNG\r\n\x1a\n"
        and "png" in png_type
        and "no-store" not in png_cache
    )
    check("GET /apple-touch-icon.png is a cacheable PNG", png_ok, f"{code} {png_type} {png_cache}")
    code, payload, _ = _http(f"{base}/")
    html = payload.decode("utf-8", errors="replace") if isinstance(payload, (bytes, bytearray)) else str(payload)
    check(
        "tab icons are root PNG/ICO for Safari",
        'href="/favicon.ico"' in html
        and 'href="/apple-touch-icon.png"' in html
        and "image/svg+xml" not in html,
        html[html.find("<title>") : html.find("</title>") + 8] if "<title>" in html else html[:120],
    )
    code, payload, _ = _http(f"{base}/superplay")
    sp_html = payload.decode("utf-8", errors="replace") if isinstance(payload, (bytes, bytearray)) else str(payload)
    check(
        "GET /superplay is the client desk",
        code == 200 and "hours-host" in sp_html and "SuperPlay — Lisenbart Board" in sp_html,
        sp_html[sp_html.find("<title>") : sp_html.find("</title>") + 8] if "<title>" in sp_html else str(code),
    )
    try:
        with urllib.request.urlopen(f"{base}/hours", timeout=2) as res:
            check(
                "GET /hours lands on /superplay",
                res.status == 200 and "/superplay" in (res.geturl() or ""),
                res.geturl(),
            )
    except urllib.error.HTTPError as e:
        loc = (e.headers.get("Location") if e.headers else "") or ""
        check(
            "GET /hours redirects to /superplay",
            e.code in {301, 302, 303, 307, 308} and "/superplay" in loc,
            f"{e.code} {loc}",
        )

    t0 = time.monotonic()
    code, payload, _ = _http(f"{base}/api/roster", timeout=2)
    dt = time.monotonic() - t0
    roles = [r.get("role") for r in (payload.get("studio") or [])]
    check("GET /api/roster under 0.4s", code == 200 and dt < 0.4, f"status={code} {dt:.3f}s")
    check(
        "roster seats are DL LP TL MD AD DEV",
        roles == ["finance", "manager", "teamlead", "designer", "art_director", "programmer"],
        str(roles),
    )

    code, payload, _ = _http(f"{base}/api/login", {"door": "studio", "role": "finance", "pin": "0000"})
    check("wrong PIN is 403", code == 403, str(payload))

    code, payload, _ = _http(f"{base}/api/login", {"door": "studio", "role": "finance", "pin": "1726"})
    check(
        "LP PIN on DL door is rejected",
        code == 403 and "cannot open" in str(payload.get("error") or ""),
        str(payload),
    )

    code, payload, _ = _http(f"{base}/api/login", {"door": "studio", "role": "manager", "pin": "2026"})
    check(
        "TL PIN on LP door is rejected",
        code == 403 and "cannot open" in str(payload.get("error") or ""),
        str(payload),
    )

    code, payload, _ = _http(f"{base}/api/login", {"door": "studio", "role": "designer", "pin": "1726"})
    team_roles = {s.get("role") for s in (payload.get("seats") or [])}
    check(
        "LP PIN on MD door sits MD",
        code == 200 and payload.get("needSeat") is True and team_roles == {"designer"},
        str(payload.get("seats") or payload),
    )

    code, payload, _ = _http(f"{base}/api/login", {"door": "studio", "role": "designer", "pin": "3245"})
    team_roles = {s.get("role") for s in (payload.get("seats") or [])}
    check(
        "DL PIN on MD door sits MD",
        code == 200 and payload.get("needSeat") is True and team_roles == {"designer"},
        str(payload.get("seats") or payload),
    )

    code, payload, _ = _http(f"{base}/api/login", {"door": "studio", "role": "manager", "pin": "3245"})
    sess = payload.get("session") or {}
    check(
        "DL PIN on LP door sits LP",
        code == 200 and payload.get("needSeat") is not True and sess.get("role") == "manager",
        str(payload.get("error") or sess),
    )

    code, payload, _ = _http(f"{base}/api/login", {"door": "studio", "role": "designer", "pin": "4429"})
    team_roles = {s.get("role") for s in (payload.get("seats") or [])}
    check(
        "team PIN on MD asks which designer",
        code == 200 and payload.get("needSeat") is True and team_roles == {"designer"},
        str(payload.get("seats") or payload),
    )

    code, payload, _ = _http(f"{base}/api/login", b"not-json")
    check("bad JSON is 400", code == 400, str(payload))

    code, payload, _ = _http(f"{base}/api/login", [])
    check("array JSON is 400", code == 400 and "object" in str(payload.get("error") or payload), str(payload))

    code, payload, headers = _http(
        f"{base}/api/login",
        {"door": "studio", "role": "finance", "pin": "3245"},
    )
    cookie = headers.get("Set-Cookie") if hasattr(headers, "get") else None
    sess = payload.get("session") or {}
    check(
        "finance login",
        code == 200 and payload.get("ok") is True and sess.get("role") == "finance",
        str(payload.get("error") or sess),
    )
    check("login sets cookie", bool(cookie), str(cookie))
    if cookie:
        ck = cookie.split(";", 1)[0]
        with server.LOCK:
            sample = next((t for t in (server.peek_db().get("tasks") or []) if t.get("id")), None)
        tid = str((sample or {}).get("id") or "t-presence")
        code, payload, _ = _http(f"{base}/api/presence", {"id": tid}, cookie=ck)
        rows = ((payload or {}).get("presence") or {}).get(tid) or []
        check(
            "DL presence marks the open card",
            code == 200 and any(r.get("self") for r in rows),
            str(payload),
        )
        code, payload, _ = _http(f"{base}/api/presence", {"id": ""}, cookie=ck)
        check("DL presence clears the card", code == 200 and tid not in ((payload or {}).get("presence") or {}), str(payload))

    code, payload, headers = _http(f"{base}/api/login", {"door": "studio", "role": "manager", "pin": "1726"})
    lp_cookie = headers.get("Set-Cookie") if hasattr(headers, "get") else None
    lp_sess = payload.get("session") or {}
    check(
        "LP login for playable crew save",
        code == 200 and lp_sess.get("role") == "manager" and bool(lp_cookie),
        str(payload.get("error") or lp_sess),
    )
    if lp_cookie:
        lp_ck = lp_cookie.split(";", 1)[0]
        stamp = datetime.utcnow().strftime("%H%M%S")
        play_name = f"DS-P-TEST-CREW-{stamp}"
        code, payload, _ = _http(
            f"{base}/api/tasks/create",
            {"name": play_name, "month": "2026-10", "role": "manager", "actorId": "manager", "direction": "playables"},
            cookie=lp_ck,
        )
        created = (payload or {}).get("task") or {}
        play_id = str(created.get("id") or "")
        check(
            "LP creates Playables task",
            code == 200 and play_id.startswith("t-") and created.get("studioStatus") == "new",
            str(payload.get("error") or created),
        )
        tiny = (
            "data:image/png;base64,"
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/comment",
            {
                "id": play_id,
                "body": "logo left",
                "attachments": [{"mime": "image/png", "data": tiny}],
                "actorId": "manager",
                "role": "manager",
            },
            cookie=lp_ck,
        )
        comment = (payload or {}).get("comment") or {}
        check(
            "LP screenshot note lands on the card",
            code == 200
            and comment.get("body") == "logo left"
            and any((a.get("kind") == "image") for a in comment.get("attachments") or []),
            str(payload.get("error") or comment) if isinstance(payload, dict) else str(payload),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": play_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "new",
                "assigneeId": "artdir1",
                "assigneeHours": 10,
                "programmerId": "programmer1",
                "programmerHours": 25,
                "help": [],
                "programmerHelp": [],
            },
            cookie=lp_ck,
        )
        saved = (payload or {}).get("task") or {}
        check(
            "LP save keeps AD, DEV, and crew hours",
            code == 200
            and saved.get("assigneeId") == "artdir1"
            and saved.get("programmerId") == "programmer1"
            and abs(float(saved.get("assigneeHours") or 0) - 10) < 0.05
            and abs(float(saved.get("programmerHours") or 0) - 25) < 0.05
            and abs(float(saved.get("producerHours") or 0) - 35) < 0.05,
            str(payload.get("error") or {k: saved.get(k) for k in ("assigneeId", "assigneeHours", "programmerId", "programmerHours", "producerHours")}),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": play_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "wip",
                "assigneeId": "artdir1",
                "assigneeHours": 10,
                "programmerId": "programmer1",
                "programmerHours": 25,
            },
            cookie=lp_ck,
        )
        moved = (payload or {}).get("task") or {}
        check(
            "LP can set Playables WIP without AD",
            code == 200 and moved.get("studioStatus") == "wip",
            str(payload.get("error") or moved.get("studioStatus")),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": play_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "artready",
                "assigneeId": "artdir1",
                "assigneeHours": 10,
                "programmerId": "programmer1",
                "programmerHours": 25,
            },
            cookie=lp_ck,
        )
        ready = (payload or {}).get("task") or {}
        check(
            "LP can set Playables Art ready without AD",
            code == 200 and ready.get("studioStatus") == "artready" and ready.get("playableHtml") is True,
            str(payload.get("error") or {k: ready.get(k) for k in ("studioStatus", "playableHtml")}),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": play_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "new",
                "assigneeId": "artdir1",
                "assigneeHours": 10,
                "programmerId": "programmer1",
                "programmerHours": 25,
            },
            cookie=lp_ck,
        )
        reset = (payload or {}).get("task") or {}
        check(
            "LP Playables status reset after cover",
            code == 200 and reset.get("studioStatus") == "new",
            str(payload.get("error") or reset.get("studioStatus")),
        )
        code, payload, _ = _http(f"{base}/api/me", cookie=lp_ck)
        reopened = next(
            (
                t
                for t in ((payload.get("state") or {}).get("tasks") or [])
                if str(t.get("id") or "") == play_id
            ),
            None,
        )
        check(
            "reopen keeps AD and DEV on the card",
            bool(reopened)
            and reopened.get("assigneeId") == "artdir1"
            and reopened.get("programmerId") == "programmer1"
            and abs(float(reopened.get("assigneeHours") or 0) - 10) < 0.05
            and abs(float(reopened.get("programmerHours") or 0) - 25) < 0.05,
            str({k: (reopened or {}).get(k) for k in ("assigneeId", "assigneeHours", "programmerId", "programmerHours")}),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": play_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "new",
                "assigneeId": "artdir1",
                "assigneeHours": 8,
                "help": [{"id": "artdir2", "hours": 2}],
                "programmerId": "programmer1",
                "programmerHours": 25,
                "programmerHelp": [],
            },
            cookie=lp_ck,
        )
        split_saved = (payload or {}).get("task") or {}
        split_help = {str((row or {}).get("id") or ""): (row or {}).get("hours") for row in (split_saved.get("help") or [])}
        check(
            "LP save keeps Playables split person and subtracts from owner",
            code == 200
            and split_saved.get("assigneeId") == "artdir1"
            and abs(float(split_saved.get("assigneeHours") or 0) - 8) < 0.05
            and abs(float(split_help.get("artdir2") or 0) - 2) < 0.05
            and abs(float(split_saved.get("producerHours") or 0) - 35) < 0.05,
            str(
                payload.get("error")
                or {
                    k: split_saved.get(k)
                    for k in ("assigneeId", "assigneeHours", "help", "programmerId", "programmerHours", "producerHours")
                }
            ),
        )
        code, payload, _ = _http(f"{base}/api/me", cookie=lp_ck)
        split_reopen = next(
            (
                t
                for t in ((payload.get("state") or {}).get("tasks") or [])
                if str(t.get("id") or "") == play_id
            ),
            None,
        )
        reopen_help = {
            str((row or {}).get("id") or ""): (row or {}).get("hours") for row in ((split_reopen or {}).get("help") or [])
        }
        check(
            "reopen keeps Playables split person",
            bool(split_reopen)
            and split_reopen.get("assigneeId") == "artdir1"
            and abs(float(split_reopen.get("assigneeHours") or 0) - 8) < 0.05
            and abs(float(reopen_help.get("artdir2") or 0) - 2) < 0.05,
            str(
                {k: (split_reopen or {}).get(k) for k in ("assigneeId", "assigneeHours", "help")}
                if split_reopen
                else payload.get("error")
            ),
        )
        _http(
            f"{base}/api/tasks/delete",
            {"id": play_id, "role": "manager", "actorId": "manager"},
            cookie=lp_ck,
        )
        stamp = datetime.utcnow().strftime("%H%M%S")
        urgent_name = f"DD-TEST-URGENT-{stamp}"
        code, payload, _ = _http(
            f"{base}/api/tasks/create",
            {"name": urgent_name, "month": "2026-10", "role": "manager", "actorId": "manager"},
            cookie=lp_ck,
        )
        created = (payload or {}).get("task") or {}
        urgent_id = str(created.get("id") or "")
        check(
            "LP creates MGX task for urgent flag",
            code == 200 and urgent_id.startswith("t-") and created.get("urgent") is False,
            str(payload.get("error") or created),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": urgent_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "wip",
            },
            cookie=lp_ck,
        )
        moved = (payload or {}).get("task") or {}
        check(
            "LP can set MGX WIP",
            code == 200 and moved.get("studioStatus") == "wip",
            str(payload.get("error") or moved.get("studioStatus")),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": urgent_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "done",
            },
            cookie=lp_ck,
        )
        check(
            "LP MGX Done still needs a result link",
            code == 400 and "result" in str(payload.get("error") or "").lower(),
            str(payload),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": urgent_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "done",
                "resultUrl": "https://board.example/mgx-done",
            },
            cookie=lp_ck,
        )
        ready = (payload or {}).get("task") or {}
        check(
            "LP can set MGX Done without TL Check",
            code == 200 and ready.get("studioStatus") == "done",
            str(payload.get("error") or ready.get("studioStatus")),
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": urgent_id,
                "role": "manager",
                "actorId": "manager",
                "urgent": True,
            },
            cookie=lp_ck,
        )
        flagged = (payload or {}).get("task") or {}
        check(
            "LP can mark a task URGENT",
            code == 200 and flagged.get("urgent") is True,
            str(payload.get("error") or {k: flagged.get(k) for k in ("urgent", "studioStatus")}),
        )
        dl_ck = cookie.split(";", 1)[0] if cookie else ""
        if dl_ck:
            code, payload, _ = _http(
                f"{base}/api/tasks/update",
                {
                    "id": urgent_id,
                    "role": "finance",
                    "actorId": "finance",
                    "urgent": False,
                },
                cookie=dl_ck,
            )
            cleared = (payload or {}).get("task") or {}
            check(
                "DL can clear URGENT",
                code == 200 and cleared.get("urgent") is False,
                str(payload.get("error") or cleared),
            )
        _http(
            f"{base}/api/tasks/update",
            {
                "id": urgent_id,
                "role": "manager",
                "actorId": "manager",
                "urgent": True,
            },
            cookie=lp_ck,
        )
        code, payload, _ = _http(
            f"{base}/api/tasks/update",
            {
                "id": urgent_id,
                "role": "manager",
                "actorId": "manager",
                "studioStatus": "closed",
            },
            cookie=lp_ck,
        )
        shut = (payload or {}).get("task") or {}
        check(
            "Closed clears URGENT",
            code == 200 and shut.get("studioStatus") == "closed" and shut.get("urgent") is False,
            str(payload.get("error") or {k: shut.get(k) for k in ("studioStatus", "urgent")}),
        )
        _http(
            f"{base}/api/tasks/delete",
            {"id": urgent_id, "role": "manager", "actorId": "manager"},
            cookie=lp_ck,
        )

    code, payload, _ = _http(f"{base}/api/hours/file?month=2026-10")
    check("hours xlsx requires sign in", code == 401)

    code, payload, _ = _http(f"{base}/api/login", {"door": "client", "email": "Olena-test", "pin": "0000"})
    check("SuperPlay wrong PIN is 403", code == 403, str(payload))

    code, payload, headers = _http(
        f"{base}/api/login",
        {"door": "client", "email": "Olena-test", "pin": "3500"},
    )
    sp_cookie = headers.get("Set-Cookie") if hasattr(headers, "get") else None
    sp_sess = payload.get("session") or {}
    sp_state = payload.get("state") or {}
    sp_tasks = sp_state.get("tasks") or []
    check(
        "SuperPlay login accepts a typed name",
        code == 200 and sp_sess.get("role") == "client" and bool(sp_cookie),
        str(payload.get("error") or sp_sess),
    )
    check(
        "SuperPlay hours payload has STATUS",
        bool(sp_tasks) and all("studioStatus" in t for t in sp_tasks),
        str(sp_tasks[:1] if sp_tasks else None),
    )
    check("SuperPlay hours hide the sign-in log", "superplayVisits" not in sp_state)
    check(
        "SuperPlay hours hide assignee and brief",
        not sp_tasks or ("assigneeId" not in sp_tasks[0] and "brief" not in sp_tasks[0]),
        str(sp_tasks[:1] if sp_tasks else None),
    )
    if sp_cookie:
        code, payload, _ = _http(f"{base}/api/me", cookie=sp_cookie.split(";", 1)[0])
        me_state = (payload or {}).get("state") or {}
        check(
            "SuperPlay /api/me has no visit log",
            code == 200 and "superplayVisits" not in me_state,
            str(list(me_state)[:12]),
        )
    if cookie:
        code, payload, _ = _http(f"{base}/api/me", cookie=cookie.split(";", 1)[0])
        visits = ((payload or {}).get("state") or {}).get("superplayVisits") or []
        check(
            "DL sees SuperPlay sign-in Login",
            any(v.get("login") == "Olena-test" for v in visits),
            str(visits[:3]),
        )
        check(
            "DL sign-in row has IP place",
            any(v.get("login") == "Olena-test" and (v.get("isp") or v.get("country")) for v in visits),
            str(visits[:3]),
        )

    if cookie:
        t0 = time.monotonic()
        code, payload, headers = _http(
            f"{base}/api/hours/file?month=2026-10&tz=Europe/Kyiv",
            cookie=cookie.split(";", 1)[0],
            timeout=3,
        )
        dt = time.monotonic() - t0
        check(
            "hours xlsx built on download",
            code == 200 and isinstance(payload, (bytes, bytearray)) and payload[:2] == b"PK" and dt < 0.5,
            f"status={code} {dt:.3f}s type={type(payload).__name__}",
        )
        if isinstance(payload, (bytes, bytearray)) and payload[:2] == b"PK":
            wb = load_workbook(BytesIO(payload))
            stamp = str(wb.active.cell(1, 1).value or "")
            names = [str(wb.active.cell(r, 1).value or "") for r in range(1, min(wb.active.max_row or 1, 40) + 1)]
            wb.close()
            check("hours xlsx stamp uses tz", "Kyiv" in stamp and stamp.startswith("Generated "), stamp)
            check(
                "hours xlsx has direction totals",
                any(n.startswith("MGX TOTAL") for n in names) or any(n.startswith("Playable TOTAL") for n in names),
                str(names[-6:]),
            )

        t0 = time.monotonic()
        code, payload, _ = _http(f"{base}/api/me", cookie=cookie.split(";", 1)[0], timeout=5)
        dt = time.monotonic() - t0
        check("/api/me under 0.3s", code == 200 and payload.get("session") and dt < 0.3, f"status={code} {dt:.3f}s")

        t0 = time.monotonic()
        code, payload, _ = _http(f"{base}/api/finance/snapshot?month=2026-10", cookie=cookie.split(";", 1)[0], timeout=2)
        dt = time.monotonic() - t0
        check(
            "finance snapshot under 0.3s",
            code == 200 and payload.get("hours") is not None and dt < 0.3,
            f"status={code} {dt:.3f}s err={payload.get('error') if isinstance(payload, dict) else payload}",
        )

        ck = cookie.split(";", 1)[0]
        with server.LOCK:
            oleksa = next((p for p in (server.peek_db().get("people") or []) if p.get("id") == "designer5"), None)
            before_card = (oleksa or {}).get("bankCard") or ""
            before_iban = (oleksa or {}).get("iban") or ""
        if oleksa:
            code, payload, _ = _http(
                f"{base}/api/people/update",
                {
                    "role": "finance",
                    "people": [{"id": "designer5", "bankCard": "4323 3450 3785 1582", "iban": before_iban}],
                },
                cookie=ck,
            )
            saved = next((p for p in (payload.get("people") or []) if p.get("id") == "designer5"), None)
            check(
                "finance can edit bank card",
                code == 200 and (saved or {}).get("bankCard") == "4323345037851582",
                str(saved),
            )
            restore = {"people": [{"bankCard": before_card}]}
            finance_io.apply_bank_card_fixes(restore)
            _http(
                f"{base}/api/people/update",
                {
                    "role": "finance",
                    "people": [{"id": "designer5", "bankCard": restore["people"][0]["bankCard"], "iban": before_iban}],
                },
                cookie=ck,
            )

        orig_anna_pin = str((auth_io.ensure_pins().get("people") or {}).get("artdir1") or "").strip()
        try:
            code, payload, _ = _http(
                f"{base}/api/people/update",
                {"role": "finance", "people": [{"id": "artdir1", "pin": "8811"}]},
                cookie=ck,
            )
            saved = next(
                (p for p in ((payload.get("state") or {}).get("people") or []) if p.get("id") == "artdir1"),
                None,
            )
            check(
                "finance can set a person PIN",
                code == 200 and (saved or {}).get("pin") == "8811",
                str(saved or payload),
            )
            code, login, _ = _http(
                f"{base}/api/login",
                {"door": "studio", "role": "art_director", "pin": "8811"},
            )
            sess = login.get("session") or {}
            check(
                "personal PIN logs that AD in",
                code == 200 and sess.get("role") == "art_director" and sess.get("personId") == "artdir1",
                str(sess or login),
            )
            code, old_login, _ = _http(
                f"{base}/api/login",
                {"door": "studio", "role": "art_director", "pin": "4429"},
            )
            old_sess = old_login.get("session") or {}
            old_seats = [s.get("id") for s in (old_login.get("seats") or [])]
            check(
                "old team PIN no longer sits the AD who changed it",
                old_sess.get("personId") != "artdir1" and "artdir1" not in old_seats,
                str(old_login),
            )
        finally:
            people = []
            with server.LOCK:
                people = list(server.peek_db().get("people") or [])
            if orig_anna_pin:
                auth_io.set_person_pin(people, "artdir1", orig_anna_pin)
            else:
                auth_io.drop_person_pin("artdir1")

        t0 = time.monotonic()
        code, payload, _ = _http(
            f"{base}/api/excel/save",
            {"month": "2026-10", "role": "finance"},
            cookie=cookie.split(";", 1)[0],
            timeout=2,
        )
        dt = time.monotonic() - t0
        check(
            "excel/save returns before Drive",
            code == 200 and payload.get("ok") is True and dt < 0.5,
            f"status={code} {dt:.3f}s {payload.get('error') if isinstance(payload, dict) else payload}",
        )

        hits = []

        def ping():
            t = time.monotonic()
            try:
                c, body, _ = _http(f"{base}/api/me", cookie=cookie.split(";", 1)[0], timeout=5)
                hits.append((c, time.monotonic() - t, bool((body or {}).get("session"))))
            except Exception as e:
                hits.append((0, time.monotonic() - t, str(e)))

        threads = [threading.Thread(target=ping) for _ in range(16)]
        t0 = time.monotonic()
        for th in threads:
            th.start()
        for th in threads:
            th.join(timeout=8)
        wall = time.monotonic() - t0
        ok_hits = [h for h in hits if h[0] == 200 and h[2] is True]
        slow = [h for h in hits if h[1] >= 5]
        check("16 concurrent /api/me all 200", len(ok_hits) == 16, str(hits[:3]))
        check("concurrent /api/me none hung", not slow and wall < 6, f"wall={wall:.3f}s slow={slow}")

        crash_id = "t-crash-del"
        with server.LOCK:
            db = server.load_db()
            db["tasks"] = [t for t in (db.get("tasks") or []) if t.get("id") != crash_id]
            db.setdefault("tasks", []).insert(
                0,
                {
                    "id": crash_id,
                    "name": "DD-CRASH-DEL-13H",
                    "line": "DD",
                    "month": "2026-10",
                    "studioStatus": "new",
                    "mondayStatus": "Ready to Start",
                    "shifts": 0,
                    "hours": 0,
                    "help": [],
                    "source": "portal",
                },
            )
            server.save_db(db)
        t0 = time.monotonic()
        code, payload, _ = _http(
            f"{base}/api/tasks/delete",
            {"id": crash_id, "role": "finance", "actorId": "finance"},
            cookie=cookie.split(";", 1)[0],
            timeout=3,
        )
        dt = time.monotonic() - t0
        gone = not any(
            str(t.get("id") or "") == crash_id for t in ((payload or {}).get("state") or {}).get("tasks") or []
        )
        check(
            "delete returns under 1s",
            code == 200 and payload.get("ok") is True and dt < 1,
            f"status={code} {dt:.3f}s err={payload.get('error') if isinstance(payload, dict) else payload}",
        )
        check("delete drops the card before Drive", gone, str((payload or {}).get("drive")))

        original_excel = server.sync_excel

        def slow_excel(db, month=None, calculator=False):
            time.sleep(1.0)
            return {"ok": False, "skipped": True}

        server.sync_excel = slow_excel
        try:
            threading.Thread(target=lambda: server._sync_files_bg("2026-10", False), daemon=True).start()
            time.sleep(0.05)
            times = []

            def ping_fast(path):
                t = time.monotonic()
                try:
                    c, body, _ = _http(f"{base}{path}", timeout=1)
                    times.append((path, c, time.monotonic() - t, body))
                except Exception as e:
                    times.append((path, 0, time.monotonic() - t, str(e)))

            workers = [
                threading.Thread(target=ping_fast, args=("/api/roster",)),
                threading.Thread(target=ping_fast, args=("/health",)),
            ]
            for w in workers:
                w.start()
            for w in workers:
                w.join(timeout=2)
            slow_hits = [row for row in times if row[2] >= 0.5]
            check(
                "roster/health stay instant while excel writes",
                times and not slow_hits and all(row[1] == 200 for row in times),
                str(times),
            )
        finally:
            server.sync_excel = original_excel


def test_drive_folder_lock() -> None:
    posts: list = []
    orig_find = drive_io._find_child
    orig_api = drive_io._api
    orig_created = dict(drive_io._CREATED_FOLDERS)
    drive_io._CREATED_FOLDERS.clear()

    def fake_find(token, parent_id, title, mime=None):
        fid = drive_io._CREATED_FOLDERS.get((str(parent_id or ""), str(title or "")))
        return {"id": fid} if fid else None

    def fake_api(token, method, path, query=None, body=None):
        if method == "POST":
            time.sleep(0.05)
            posts.append(body)
            return {"id": "folder-only-once"}
        return {"files": []}

    drive_io._find_child = fake_find
    drive_io._api = fake_api
    try:
        got = []

        def make():
            got.append(drive_io._ensure_folder("tok", "parent", "DS-P-TWIN"))

        threads = [threading.Thread(target=make) for _ in range(2)]
        for th in threads:
            th.start()
        for th in threads:
            th.join(timeout=2)
        check("Drive ensure_folder posts once under race", len(posts) == 1, str(len(posts)))
        check(
            "Drive ensure_folder returns the same id",
            len(got) == 2 and got[0][0] == "folder-only-once" and got[1][0] == "folder-only-once",
            str(got),
        )
        py = (ROOT / "drive_io.py").read_text(encoding="utf-8")
        check("create_live_pack serializes per name", "with _pack_lock(name):" in py)
    finally:
        drive_io._find_child = orig_find
        drive_io._api = orig_api
        drive_io._CREATED_FOLDERS.clear()
        drive_io._CREATED_FOLDERS.update(orig_created)


def start_local_server() -> tuple[threading.Thread, str]:
    import server as srv

    httpd = srv.ThreadingHTTPServer(("127.0.0.1", srv.PORT), srv.Handler)
    httpd.daemon_threads = True
    httpd.timeout = 8
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    return thread, f"http://127.0.0.1:{srv.PORT}"


def main() -> int:
    print("Excel / db")
    for fn in (
        test_client_hours_xlsx,
        test_pin_opens,
        test_person_pins,
        test_parse_past_total,
        test_apply_full_sheet,
        test_write_month_roundtrip,
        test_corrupt_xlsx_empty,
        test_live_workbook_if_present,
        test_mgx_seat_headers,
        test_corrupt_db_recovery,
        test_save_db_atomic,
        test_unreadable_ingest_does_not_wipe,
        test_sync_files_releases_lock,
        test_employeers_pay_details,
        test_playables_desk,
        test_drive_folder_lock,
        test_playable_card_pipeline,
        test_card_presence,
        test_urgent_marker,
        test_card_folder_buttons,
        test_superplay_visits_and_status,
        test_lp_can_close_without_send,
        test_status_receipt,
        test_comment_shots,
        test_finance_reads_board_not_xlsx,
    ):
        try:
            fn()
        except Exception:
            check(fn.__name__, False, traceback.format_exc()[-400:])

    print("HTTP")
    base = f"http://127.0.0.1:{server.PORT}"
    owned = False
    try:
        urllib.request.urlopen(f"{base}/health", timeout=1).read()
    except Exception:
        print("  starting local desk…")
        start_local_server()
        owned = True
        for _ in range(40):
            try:
                urllib.request.urlopen(f"{base}/health", timeout=0.25).read()
                break
            except Exception:
                time.sleep(0.1)
        else:
            check("desk starts", False, "no /health")
            print(f"\n{PASS} passed, {FAIL} failed")
            return 1
    try:
        test_http(base)
    except Exception:
        check("HTTP suite", False, traceback.format_exc()[-400:])

    print(f"\n{PASS} passed, {FAIL} failed")
    if ERRORS:
        print("\n".join(ERRORS))
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
